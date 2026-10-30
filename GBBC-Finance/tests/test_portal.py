from datetime import date
from pathlib import Path
import io
import re
import pytest
from pypdf import PdfReader
from werkzeug.security import generate_password_hash
from app import create_app, db, User, Person, Fund, FiscalYear, Gift, SignupRequest, ensure_schema

def _hash(password):
    return generate_password_hash(password, method='pbkdf2:sha256')

def pdf_text(data):
    reader = PdfReader(io.BytesIO(data))
    return '\n'.join((page.extract_text() or '') for page in reader.pages)

@pytest.fixture
def client(tmp_path):
    app = create_app({'TESTING': True, 'SECRET_KEY': 'test-key', 'SQLALCHEMY_DATABASE_URI': f'sqlite:///{tmp_path / "test.db"}'})
    with app.app_context():
        db.create_all()
        ensure_schema()
        db.session.add_all([
            User(first_name='Sally', last_name='Admin', email='sally@example.org', phone='555-1000', role='Super Admin', password_hash=_hash('correct-password')),
            User(first_name='Chris', last_name='Leader', email='chris@example.org', phone='555-2000', role='Church Leader', password_hash=_hash('correct-password')),
            Person(first_name='Ada', last_name='Member', member_type='Member'),
            FiscalYear(name='2026', start_date=date(2026,1,1), end_date=date(2026,12,31)),
        ]); db.session.commit()
    return app.test_client()

def token(client):
    with client.session_transaction() as s: return s['csrf']

def login(client, email):
    client.get('/login')
    return client.post('/login', data={'csrf_token': token(client), 'email': email, 'password': 'correct-password'}, follow_redirects=True)

def test_auth_and_user_fields(client):
    assert client.get('/people').status_code == 302
    assert b'Good morning, Sally' in login(client, 'sally@example.org').data
    bad = client.post('/admin/users', data={'csrf_token': token(client), 'first_name':'Pat','last_name':'Smith','email':'pat@example.org','role':'Finance Staff','password':'secure-password'}, follow_redirects=True)
    assert b'phone' in bad.data.lower()
    good = client.post('/admin/users', data={'csrf_token': token(client), 'first_name':'Pat','last_name':'Smith','email':'pat@example.org','phone':'555-3000','role':'Finance Staff','password':'secure-password'}, follow_redirects=True)
    assert b'User created' in good.data
    assert client.post('/admin/users', data={}).status_code == 400
    assert b'Church Leader' in client.get('/admin/roles').data
    assert b'People-Manager' in client.get('/admin/roles').data

def test_search_engines_are_blocked(client):
    robots = client.get('/robots.txt')
    assert robots.status_code == 200
    assert robots.data == b'User-agent: *\nDisallow: /\n'
    assert 'noindex' in client.get('/login').headers['X-Robots-Tag']
    assert 'noindex' in client.get('/people').headers['X-Robots-Tag']

def test_security_headers(client):
    headers = client.get('/login').headers
    assert "frame-ancestors 'none'" in headers['Content-Security-Policy']
    assert "object-src 'none'" in headers['Content-Security-Policy']
    assert 'Strict-Transport-Security' not in headers

def test_login_next_rejects_offsite_redirects(client):
    for bad in ('//evil.example', '/\\evil.example', '/\t/evil.example', 'https://evil.example'):
        client.get('/login')
        response = client.post(f'/login?next={bad}', data={'csrf_token': token(client), 'email': 'sally@example.org', 'password': 'correct-password'})
        assert response.status_code == 302
        assert 'evil.example' not in response.headers['Location']
        client.post('/logout', data={'csrf_token': token(client)})
    client.get('/login')
    response = client.post('/login?next=/people', data={'csrf_token': token(client), 'email': 'sally@example.org', 'password': 'correct-password'})
    assert response.headers['Location'].endswith('/people')

def test_people_manager_role_permissions(client):
    with client.application.app_context():
        db.session.add(User(
            first_name='Pam', last_name='Manager', email='pam@example.org', phone='555-4000',
            role='People-Manager', password_hash=_hash('correct-password'),
        ))
        db.session.commit()
    login(client, 'pam@example.org')
    assert client.get('/').status_code == 200
    assert client.get('/people').status_code == 200
    assert client.get('/people/new').status_code == 200
    assert client.get('/help').status_code == 200
    assert b'session_idle.js' in client.get('/people').data
    for path in ['/contributions', '/reports', '/utilities', '/admin/users', '/admin/roles']:
        assert client.get(path).status_code == 403
    client.post('/logout', data={'csrf_token': token(client)})
    signup = client.get('/signup').data
    assert b'People-Manager' in signup

def test_backup_and_restore_year_workbook(client):
    from decimal import Decimal
    from openpyxl import Workbook, load_workbook
    from app import Audit, _cell_amount
    assert _cell_amount(10.1) == Decimal('10.10')
    assert _cell_amount(55.25) == Decimal('55.25')
    assert _cell_amount('12.50') == Decimal('12.50')
    login(client, 'sally@example.org')
    page = client.get('/admin/backup')
    assert page.status_code == 200
    assert b'Download Excel backup' in page.data
    assert b'Restore backup' in page.data
    with client.application.app_context():
        db.session.add(Person(first_name='Ben', last_name='Backup', member_type='Member'))
        db.session.flush()
        person_id = db.session.scalar(db.select(Person.id).filter_by(last_name='Backup'))
        db.session.add(Gift(
            person_id=person_id, fiscal_year_id=1, fund_id=1, date=date(2026, 3, 1),
            amount=55.25, method='Check', note='Backup gift', entered_by_id=1,
        ))
        db.session.add(Gift(
            person_id=person_id, fiscal_year_id=1, fund_id=1, date=date(2026, 3, 2),
            amount=Decimal('10.10'), method='Cash', note='Float prone', entered_by_id=1,
        ))
        db.session.commit()
    download = client.post('/admin/backup', data={
        'csrf_token': token(client), 'action': 'download', 'year_id': '1',
    })
    assert download.status_code == 200
    assert download.mimetype.startswith('application/vnd.openxmlformats-officedocument')
    wb = load_workbook(io.BytesIO(download.data))
    assert set(wb.sheetnames) == {'People', 'Contributions', 'Users', 'Audit log'}
    gift_names = [row[2].value for row in wb['Contributions'].iter_rows(min_row=2) if row[2].value]
    assert 'Ben Backup' in gift_names
    # Simulate Excel float cells that used to break restore.
    float_wb = Workbook()
    float_wb.properties.subject = wb.properties.subject
    for title in ('People', 'Contributions', 'Users', 'Audit log'):
        if title == 'People':
            src, dst = wb[title], float_wb.active
            dst.title = title
        else:
            src, dst = wb[title], float_wb.create_sheet(title)
        for row in src.iter_rows(values_only=True):
            dst.append(list(row))
    for row in float_wb['Contributions'].iter_rows(min_row=2):
        if row[6].value not in (None, ''):
            row[6].value = float(row[6].value)
    float_buf = io.BytesIO()
    float_wb.save(float_buf)
    float_payload = float_buf.getvalue()
    with client.application.app_context():
        gift = db.session.scalar(db.select(Gift).where(Gift.note == 'Backup gift'))
        gift.amount = 1
        gift.note = 'changed'
        db.session.commit()
        gift_id = gift.id
    upload = client.post(
        '/admin/backup',
        data={
            'csrf_token': token(client),
            'action': 'upload',
            'year_id': '1',
            'backup_file': (io.BytesIO(float_payload), 'backup.xlsx'),
        },
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'Confirm restore' in upload.data
    assert b'Type RESTORE to confirm' in upload.data
    with client.session_transaction() as sess:
        pending = sess.get('backup_restore')
        assert pending and pending.get('token')
        confirm_token = pending['token']
    blocked = client.post('/admin/backup', data={
        'csrf_token': token(client),
        'action': 'confirm_restore',
        'year_id': '1',
        'confirm_token': confirm_token,
        'confirm_text': 'NOPE',
    }, follow_redirects=True)
    assert b'Restore confirmation failed' in blocked.data
    # re-upload after failed confirm cleared staging
    upload = client.post(
        '/admin/backup',
        data={
            'csrf_token': token(client),
            'action': 'upload',
            'year_id': '1',
            'backup_file': (io.BytesIO(float_payload), 'backup.xlsx'),
        },
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'Confirm restore' in upload.data
    with client.session_transaction() as sess:
        confirm_token = sess['backup_restore']['token']
    restored = client.post('/admin/backup', data={
        'csrf_token': token(client),
        'action': 'confirm_restore',
        'year_id': '1',
        'confirm_token': confirm_token,
        'confirm_text': 'RESTORE',
    }, follow_redirects=True)
    assert b'Restored backup for fiscal year' in restored.data
    assert b'positive amount' not in restored.data
    with client.application.app_context():
        gift = db.session.scalar(db.select(Gift).where(Gift.note == 'Backup gift'))
        assert gift is not None
        assert gift.amount == Decimal('55.25')
        assert db.session.get(Gift, gift_id) is None or db.session.get(Gift, gift_id).note != 'changed'
        float_gift = db.session.scalar(db.select(Gift).where(Gift.note == 'Float prone'))
        assert float_gift is not None and float_gift.amount == Decimal('10.10')
        event = db.session.scalar(db.select(Audit).where(Audit.action == 'restore').order_by(Audit.id.desc()))
        assert event is not None and '2026' in event.summary

def test_inactivity_logs_out_after_ten_minutes(client):
    from datetime import datetime, timedelta, timezone
    from app import Audit
    login(client, 'sally@example.org')
    assert client.get('/').status_code == 200
    with client.session_transaction() as sess:
        expired = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(minutes=11)
        sess['last_activity'] = expired.isoformat(timespec='seconds')
    response = client.get('/people', follow_redirects=True)
    assert response.status_code == 200
    assert b'10 minutes of inactivity' in response.data
    assert b'Sign in' in response.data or b'sign in' in response.data.lower() or b'Password' in response.data
    assert client.get('/people').status_code == 302
    assert client.get('/session/ping').status_code == 401
    with client.application.app_context():
        event = db.session.scalar(
            db.select(Audit).where(Audit.action == 'idle_logout', Audit.user_id == 1).order_by(Audit.id.desc())
        )
        assert event is not None
        assert '10 minutes of inactivity' in event.summary
    login(client, 'sally@example.org')
    audit_page = client.get('/admin/audit').data
    assert b'idle_logout' in audit_page
    assert b'Signed out after 10 minutes of inactivity' in audit_page

def test_leader_suppression_and_server_denial(client):
    with client.application.app_context():
        db.session.add(Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,9,1), amount=123, method='Check', entered_by_id=1)); db.session.commit()
    login(client, 'chris@example.org')
    dashboard = client.get('/').data
    assert b'$123.00' in dashboard and b'Ada Member' not in dashboard
    assert b'Giving by fund' in dashboard and b'General Fund' in dashboard
    assert b'Contributions</span>' not in dashboard and b'Administration</span>' not in dashboard
    for path in ['/contributions', '/reports/export.xlsx', '/admin/users', '/admin/roles', '/people/new']:
        assert client.get(path).status_code == 403
    report = client.get('/reports').data
    assert b'$123.00' in report and b'Ada Member' not in report
    assert client.get('/reports?person_id=1').status_code == 403
    assert client.get('/reports?minimum_total=1').status_code == 403

def test_dashboard_fund_pie_chart(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        general = db.session.scalar(db.select(Fund).filter_by(name='General Fund'))
        missions = db.session.scalar(db.select(Fund).filter_by(name='Missions Fund'))
        db.session.add_all([
            Gift(person_id=1, fiscal_year_id=1, fund_id=general.id, date=date(2026, 9, 1), amount=75, method='Check', entered_by_id=1),
            Gift(person_id=1, fiscal_year_id=1, fund_id=missions.id, date=date(2026, 9, 2), amount=25, method='Cash', entered_by_id=1),
            Gift(person_id=1, fiscal_year_id=1, fund_id=general.id, date=date(2026, 9, 3), amount=10, method='Cash', entered_by_id=1, status='void'),
        ])
        db.session.commit()
    page = client.get('/').data
    assert b'Giving by fund' in page
    assert b'fund-pie' in page
    assert b'General Fund' in page
    assert b'Missions Fund' in page
    assert b'$75.00' in page
    assert b'$25.00' in page
    assert b'75.0% of year-to-date giving' in page
    assert b'25.0% of year-to-date giving' in page
    assert b'Monthly giving' in page
    assert b'month-bars' in page

def test_dashboard_monthly_fund_bars(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        general = db.session.scalar(db.select(Fund).filter_by(name='General Fund'))
        missions = db.session.scalar(db.select(Fund).filter_by(name='Missions Fund'))
        building = db.session.scalar(db.select(Fund).filter_by(name='Building Fund'))
        db.session.add_all([
            Gift(person_id=1, fiscal_year_id=1, fund_id=general.id, date=date(2026, 1, 15), amount=40, method='Check', entered_by_id=1),
            Gift(person_id=1, fiscal_year_id=1, fund_id=missions.id, date=date(2026, 1, 20), amount=10, method='Cash', entered_by_id=1),
            Gift(person_id=1, fiscal_year_id=1, fund_id=building.id, date=date(2026, 3, 4), amount=30, method='Online', entered_by_id=1),
            Gift(person_id=1, fiscal_year_id=1, fund_id=general.id, date=date(2026, 3, 5), amount=5, method='Cash', entered_by_id=1, status='void'),
        ])
        db.session.commit()
    page = client.get('/').data.decode()
    assert 'Monthly giving' in page
    assert "The amount above it is that month's total, stacked by fund." in page
    assert page.count('class="month-bar"') == 12
    assert page.count('class="month-bar-total"') == 12
    assert 'class="month-bar-total">$50.00' in page
    assert 'class="month-bar-total">$30.00' in page
    assert 'class="month-bar-total">$0.00' in page
    assert 'January 2026 total $50.00' in page
    assert 'January 2026 · $50.00' in page
    assert 'General Fund · $40.00' in page
    assert 'Missions Fund · $10.00' in page
    assert 'March 2026 · $30.00' in page
    assert 'Building Fund · $30.00' in page
    assert '#1b624c' in page
    assert '#3d9b74' in page
    assert '#c4a35a' in page
    assert 'February 2026 · $0.00' in page

def test_gift_entry_and_year_close(client):
    login(client, 'sally@example.org')
    form = client.get('/contributions/new').data.decode()
    assert form.index('General Fund') < form.index('Missions Fund') < form.index('Building Fund') < form.index('Benevolence Fund')
    assert 'selected' in form[form.index('General Fund')-80:form.index('General Fund')]
    data={'csrf_token':token(client),'person_id':'1','fiscal_year_id':'1','fund_id':'1','date':'2026-09-01','amount':'23.45','method':'Check'}
    assert b'Contribution saved' in client.post('/contributions/new',data=data,follow_redirects=True).data
    with client.application.app_context():
        db.session.get(FiscalYear,1).is_open=False;db.session.commit()
    assert b'open fiscal year' in client.post('/contributions/new',data=data,follow_redirects=True).data

def test_person_structured_address_and_letter(client):
    login(client, 'sally@example.org')
    data = {
        'csrf_token': token(client),
        'first_name': 'Ada', 'last_name': 'Member', 'member_type': 'Member',
        'email': 'ada@example.org', 'phone': '555-0100',
        'address_line1': '170 Greenbrook Rd.', 'address_line2': 'Attn: Finance',
        'city': 'Green Brook', 'state': 'nj', 'zip_code': '08812',
        'audit_reason': 'Corrected mailing address',
    }
    assert b'Person saved' in client.post('/people/1/edit', data=data, follow_redirects=True).data
    with client.application.app_context():
        person = db.session.get(Person, 1)
        assert person.address_line1 == '170 Greenbrook Rd.'
        assert person.address_line2 == 'Attn: Finance'
        assert person.city == 'Green Brook'
        assert person.state == 'NJ'
        assert person.zip_code == '08812'
        assert person.address_display == '170 Greenbrook Rd., Attn: Finance, Green Brook, NJ 08812'
    with client.application.app_context():
        db.session.add(Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,9,1), amount=50, method='Check', entered_by_id=1)); db.session.commit()
    pdf = client.get('/reports/statements/1.pdf?year_id=1')
    assert pdf.status_code == 200
    letter = pdf_text(pdf.data)
    assert '170 Greenbrook Rd.' in letter
    assert 'Attn: Finance' in letter
    assert 'Green Brook, NJ 08812' in letter
    people = client.get('/people').data
    assert b'170 Greenbrook Rd.' in people and b'Green Brook, NJ 08812' in people
    form = client.get('/people/new').data
    assert b'Address line 1' in form and b'us-states' in form
    template = client.get('/people/template.csv').data
    assert b'address_line1,address_line2,city,state,zip_code' in template

def test_exports_and_role_matrix(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        db.session.add(Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,9,1), amount=123, method='Check', entered_by_id=1)); db.session.commit()
    pdf = client.get('/reports/statements/1.pdf')
    assert pdf.status_code == 200 and pdf.data.startswith(b'%PDF')
    letter = pdf_text(pdf.data)
    assert 'Romella Seepaul' in letter
    assert 'Dear Ada' in letter
    assert 'January 31, 2027' in letter
    assert letter.index('January 31, 2027') < letter.index('Ada Member')
    assert '2 Corinthians 9:6-8' in letter
    assert 'The point is this' in letter
    assert 'God loves a cheerful giver' in letter
    assert 'Itemized contributions' in letter
    # Letterhead is stamped as artwork (logo, church name, tax ID, website, green banner).
    reader = PdfReader(io.BytesIO(pdf.data))
    letter_page = reader.pages[0]
    xobjects = letter_page['/Resources']['/XObject'].get_object()
    images = [xobjects[name] for name in xobjects if xobjects[name].get('/Subtype') == '/Image']
    assert any(img.get('/Width', 0) >= 1000 and img.get('/Height', 0) >= 200 for img in images)
    # Website hyperlink matches the visible www.greenbrookchurch.org label.
    annots = letter_page.get('/Annots') or []
    uris = []
    for ann in annots:
        action = ann.get_object().get('/A')
        if action:
            uri = action.get_object().get('/URI')
            if uri:
                uris.append(str(uri))
    assert any('www.greenbrookchurch.org' in uri for uri in uris)
    # Itemized pages also carry the green letterhead artwork.
    assert len(reader.pages) >= 2
    itemized = reader.pages[1]
    item_xobjects = itemized['/Resources']['/XObject'].get_object()
    item_images = [item_xobjects[name] for name in item_xobjects if item_xobjects[name].get('/Subtype') == '/Image']
    assert any(img.get('/Width', 0) >= 1000 and img.get('/Height', 0) >= 200 for img in item_images)
    xlsx = client.get('/reports/export.xlsx')
    assert xlsx.status_code == 200 and xlsx.data.startswith(b'PK')
    assert client.get('/admin/settings').status_code == 200

def test_login_limit(client):
    client.get('/login')
    for _ in range(5):
        response = client.post('/login', data={'csrf_token':token(client), 'email':'sally@example.org', 'password':'wrong'})
        assert response.status_code == 200
    assert client.post('/login', data={'csrf_token':token(client), 'email':'sally@example.org', 'password':'correct-password'}).status_code == 429

def test_all_primary_pages_render(client):
    login(client, 'sally@example.org')
    paths = ['/', '/people', '/people/new', '/people/1/edit', '/people/import', '/contributions', '/contributions/import', '/contributions/new', '/reports', '/admin/users', '/admin/users/1/edit', '/admin/roles', '/admin/fiscal', '/admin/settings', '/admin/feedback', '/admin/audit', '/admin/backup', '/help']
    for path in paths:
        response = client.get(path)
        assert response.status_code == 200, path
        assert b'GBBC Finance' in response.data, path
    assert client.get('/approvals').status_code == 302

def test_signup_and_admin_approval_flow(client):
    login_page = client.get('/login')
    assert b'Sign Up' in login_page.data
    assert b'href="/signup"' in login_page.data
    signup_page = client.get('/signup')
    assert signup_page.status_code == 200
    assert b'Create an account' in signup_page.data
    weak = client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Pat', 'last_name': 'Applicant', 'email': 'pat@example.org',
        'password': 'short1!', 'confirm_password': 'short1!', 'role': 'People Viewer',
    }, follow_redirects=True)
    assert weak.status_code == 400
    assert b'at least 8 characters' in weak.data
    mismatch = client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Pat', 'last_name': 'Applicant', 'email': 'pat@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure2!', 'role': 'Finance Staff',
    }, follow_redirects=True)
    assert b'do not match' in mismatch.data
    signup_page = client.get('/signup')
    assert b'password-toggle' in signup_page.data
    assert b'Show' in signup_page.data
    assert b'Finance Staff' in signup_page.data
    assert b'System Administrator' in signup_page.data
    assert b'value="Admin"' not in signup_page.data
    assert b'value="Finance"' not in signup_page.data  # exact Finance role removed; Finance Staff remains
    ok = client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Pat', 'last_name': 'Applicant', 'email': 'pat@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure1!', 'role': 'People Viewer',
    }, follow_redirects=True)
    assert b'approved by an administrator' in ok.data
    assert b"administrator's email address" in ok.data
    assert b'Email delivery is not configured' in ok.data
    with client.application.app_context():
        pending = db.session.scalar(db.select(SignupRequest).filter_by(email='pat@example.org'))
        assert pending is not None and pending.status == 'pending'
        assert db.session.scalar(db.select(User).filter_by(email='pat@example.org')) is None
        outbox = Path(client.application.instance_path) / 'mail_outbox'
        received = list(outbox.glob(f'signup_{pending.id}_received_*.txt'))
        assert received, 'expected immediate signup acknowledgment email in mail_outbox'
        assert 'signup request received' in received[0].read_text(encoding='utf-8').lower()
    blocked = client.post('/login', data={'csrf_token': token(client), 'email': 'pat@example.org', 'password': 'Secure1!'}, follow_redirects=True)
    assert b'waiting for administrator approval' in blocked.data
    login(client, 'sally@example.org')
    users_page = client.get('/admin/users')
    assert users_page.status_code == 200
    assert b'Pat Applicant' in users_page.data
    assert b'Pending signup requests' in users_page.data
    assert b'User Management</span>' not in client.get('/people').data
    approved = client.post('/admin/users', data={
        'csrf_token': token(client), 'signup_id': '1', 'action': 'approve', 'audit_reason': 'Verified church member',
    }, follow_redirects=True)
    assert b'Approved Pat Applicant' in approved.data
    client.post('/logout', data={'csrf_token': token(client)})
    client.get('/login')
    signed_in = client.post('/login', data={'csrf_token': token(client), 'email': 'pat@example.org', 'password': 'Secure1!'}, follow_redirects=True)
    assert signed_in.status_code == 200
    assert b'People' in signed_in.data
    assert client.get('/').status_code == 403
    assert client.get('/people').status_code == 200
    assert client.get('/help').status_code == 200
    assert client.get('/contributions').status_code == 403
    assert client.get('/admin/users').status_code == 403
    assert b'Overview</span>' not in client.get('/people').data
    assert b'Administration</span>' not in client.get('/people').data

def test_finance_staff_signup_role_cannot_access_users_admin(client):
    client.get('/signup')
    client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Fin', 'last_name': 'Staffish', 'email': 'fin@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure1!', 'role': 'Finance Staff',
    })
    login(client, 'sally@example.org')
    with client.application.app_context():
        req = db.session.scalar(db.select(SignupRequest).filter_by(email='fin@example.org'))
        req_id = req.id
    client.post('/admin/users', data={
        'csrf_token': token(client), 'signup_id': str(req_id), 'action': 'approve', 'audit_reason': 'Finance hire',
    }, follow_redirects=True)
    client.post('/logout', data={'csrf_token': token(client)})
    client.get('/login')
    login_resp = client.post('/login', data={'csrf_token': token(client), 'email': 'fin@example.org', 'password': 'Secure1!'}, follow_redirects=True)
    assert b'Overview' in login_resp.data
    assert client.get('/').status_code == 200
    assert client.get('/people').status_code == 200
    assert client.get('/contributions').status_code == 200
    assert client.get('/reports').status_code == 200
    reports = client.get('/reports').data
    assert b'Giving statements' in reports
    assert b'Ada Member' in reports
    assert b'/reports/statements/1.pdf' in reports
    assert client.get('/reports/statements/1.pdf').status_code == 200
    assert client.get('/utilities').status_code == 200
    assert client.get('/help').status_code == 200
    assert client.get('/admin/users').status_code == 403
    home = client.get('/').data
    assert b'Reports</span>' in home
    assert b'Utilities</span>' in home
    assert b'Administration</span>' not in home

def test_system_administrator_signup_role_can_approve(client):
    client.get('/signup')
    client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Ada', 'last_name': 'Adminish', 'email': 'ada.admin@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure1!', 'role': 'System Administrator',
    })
    client.get('/signup')
    client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Vie', 'last_name': 'Wer', 'email': 'viewer2@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure1!', 'role': 'People Viewer',
    })
    login(client, 'sally@example.org')
    with client.application.app_context():
        admin_req = db.session.scalar(db.select(SignupRequest).filter_by(email='ada.admin@example.org'))
        viewer_req = db.session.scalar(db.select(SignupRequest).filter_by(email='viewer2@example.org'))
        admin_id, viewer_id = admin_req.id, viewer_req.id
    client.post('/admin/users', data={
        'csrf_token': token(client), 'signup_id': str(admin_id), 'action': 'approve', 'audit_reason': 'IT staff',
    }, follow_redirects=True)
    client.post('/logout', data={'csrf_token': token(client)})
    client.get('/login')
    client.post('/login', data={'csrf_token': token(client), 'email': 'ada.admin@example.org', 'password': 'Secure1!'}, follow_redirects=True)
    page = client.get('/admin/users')
    assert page.status_code == 200
    assert b'Pending signup requests' in page.data
    assert b'Sally Admin' in page.data
    assert b'password_hash' not in page.data
    denied = client.post('/admin/users', data={
        'csrf_token': token(client), 'signup_id': str(viewer_id), 'action': 'deny', 'audit_reason': 'Not needed',
    }, follow_redirects=True)
    assert b'Denied Vie Wer' in denied.data
    with client.application.app_context():
        assert db.session.scalar(db.select(SignupRequest).filter_by(email='viewer2@example.org')).status == 'denied'
        assert db.session.scalar(db.select(User).filter_by(email='viewer2@example.org')) is None

def test_legacy_admin_finance_roles_are_migrated(client):
    with client.application.app_context():
        db.session.add(User(
            first_name='Legacy', last_name='Admin', email='legacy.admin@example.org', phone='555-0102',
            role='Admin', password_hash=_hash('correct-password'),
        ))
        db.session.add(User(
            first_name='Legacy', last_name='Finance', email='legacy.finance@example.org', phone='555-0103',
            role='Finance', password_hash=_hash('correct-password'),
        ))
        db.session.add(SignupRequest(
            first_name='Pend', last_name='Ing', email='pending.legacy@example.org',
            password_hash=_hash('Secure1!'), requested_role='Admin', status='pending',
        ))
        db.session.commit()
        from app import migrate_legacy_roles
        migrate_legacy_roles()
        assert db.session.scalar(db.select(User).filter_by(email='legacy.admin@example.org')).role == 'System Administrator'
        assert db.session.scalar(db.select(User).filter_by(email='legacy.finance@example.org')).role == 'Finance Staff'
        assert db.session.scalar(db.select(SignupRequest).filter_by(email='pending.legacy@example.org')).requested_role == 'System Administrator'
        assert 'Admin' not in __import__('app', fromlist=['ROLES']).ROLES
        assert 'Finance' not in __import__('app', fromlist=['ROLES']).ROLES

def test_signup_sends_email_when_smtp_configured(client, monkeypatch):
    sent = []

    class FakeSMTP:
        def __init__(self, host, port, timeout=20):
            self.host, self.port = host, port
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def ehlo(self):
            sent.append('ehlo')
            return (250, b'ok')
        def starttls(self):
            sent.append('starttls')
            return (220, b'ready')
        def login(self, username, password):
            sent.append(('login', username, password))
        def send_message(self, message):
            sent.append(('send', message['To'], message['Subject'], message.get_content()))

    monkeypatch.setattr('smtplib.SMTP', FakeSMTP)
    with client.application.app_context():
        from app import set_setting
        set_setting('smtp_host', 'smtp.example.com')
        set_setting('smtp_port', '587')
        set_setting('smtp_username', 'mailer@example.com')
        set_setting('smtp_password', 'secret-pass')
        set_setting('smtp_encryption', 'starttls')
        set_setting('mail_from', 'noreply@example.com')
        db.session.commit()
    client.get('/signup')
    response = client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Mail', 'last_name': 'Tester', 'email': 'mail.tester@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure1!', 'role': 'People Viewer',
    }, follow_redirects=True)
    assert response.status_code == 200
    assert b'confirmation email was sent' in response.data
    assert any(item[0] == 'send' for item in sent if isinstance(item, tuple))
    send = next(item for item in sent if isinstance(item, tuple) and item[0] == 'send')
    assert send[1] == 'mail.tester@example.org'
    assert 'signup request received' in send[2].lower()

def test_admin_users_edit_requires_reason_and_logs_changes(client):
    client.get('/signup')
    client.post('/signup', data={
        'csrf_token': token(client),
        'first_name': 'Vie', 'last_name': 'Wer', 'email': 'viewer3@example.org',
        'password': 'Secure1!', 'confirm_password': 'Secure1!', 'role': 'People Viewer',
    })
    login(client, 'sally@example.org')
    with client.application.app_context():
        req = db.session.scalar(db.select(SignupRequest).filter_by(email='viewer3@example.org'))
        req_id = req.id
    client.post('/admin/users', data={
        'csrf_token': token(client), 'signup_id': str(req_id), 'action': 'approve', 'audit_reason': 'Approved viewer',
    }, follow_redirects=True)
    with client.application.app_context():
        account = db.session.scalar(db.select(User).filter_by(email='viewer3@example.org'))
        account_id = account.id
        password_hash = account.password_hash
    edit_page = client.get(f'/admin/users/{account_id}/edit')
    assert edit_page.status_code == 200
    assert b'Edit user' in edit_page.data
    assert b'Reason for this change' in edit_page.data
    assert b'password_hash' not in edit_page.data
    assert password_hash.encode() not in edit_page.data
    missing_reason = client.post(f'/admin/users/{account_id}/edit', data={
        'csrf_token': token(client),
        'first_name': 'View', 'last_name': 'Erson', 'email': 'viewer3@example.org',
        'username': 'viewer3', 'phone': '555-0199', 'role': 'People Viewer', 'password': '',
    }, follow_redirects=True)
    assert b'reason for this change is required' in missing_reason.data.lower()
    saved = client.post(f'/admin/users/{account_id}/edit', data={
        'csrf_token': token(client),
        'first_name': 'View', 'last_name': 'Erson', 'email': 'viewer3@example.org',
        'username': 'viewer3', 'phone': '555-0199', 'role': 'People Viewer', 'password': '',
        'audit_reason': 'Corrected display name',
    }, follow_redirects=True)
    assert b'User updated' in saved.data
    assert b'View Erson' in saved.data
    audit_page = client.get('/admin/audit')
    assert audit_page.status_code == 200
    assert b'Time (Eastern)' in audit_page.data
    assert b'Corrected display name' in audit_page.data
    assert 'First name: Vie → View'.encode() in audit_page.data
    with client.application.app_context():
        account = db.session.get(User, account_id)
        assert account.first_name == 'View' and account.phone == '555-0199' and account.username == 'viewer3'

def test_contribution_create_appears_in_audit_with_details(client):
    login(client, 'sally@example.org')
    response = client.post('/contributions/new', data={
        'csrf_token': token(client),
        'person_id': '1', 'fiscal_year_id': '1', 'date': '2026-09-10',
        'amount': '42.50', 'fund_id': '1', 'method': 'Check', 'reference': '1001', 'note': 'Sunday gift',
    }, follow_redirects=True)
    assert b'Contribution saved' in response.data
    audit_page = client.get('/admin/audit').data
    assert b'Ada Member' in audit_page
    assert b'General Fund' in audit_page
    assert b'42.50' in audit_page
    assert b'Sunday gift' in audit_page

def test_people_search_matches_any_name_part(client):
    login(client, 'sally@example.org')
    form = client.get('/contributions/new')
    assert form.status_code == 200
    assert b'person_search' in form.data
    assert b'person_search.js' in form.data
    assert b'Type at least 2 letters' in form.data
    assert client.get('/people/search?q=a').status_code == 200
    assert client.get('/people/search?q=a').get_json() == []
    with client.application.app_context():
        db.session.add(Person(first_name='Maria C', last_name='Pagulayan', member_type='Member'))
        db.session.add(Person(first_name='Jose', last_name='Garcia-Lopez', member_type='Non-member'))
        db.session.add(Person(first_name='Inactive', last_name='Person', member_type='Member', active=False))
        db.session.commit()
    by_first = client.get('/people/search?q=ria').get_json()
    assert any(row['name'] == 'Maria C Pagulayan' for row in by_first)
    by_middle = client.get('/people/search?q=C Pag').get_json()
    assert any(row['name'] == 'Maria C Pagulayan' for row in by_middle)
    by_last = client.get('/people/search?q=garcia').get_json()
    assert any(row['last_name'] == 'Garcia-Lopez' for row in by_last)
    by_tokens = client.get('/people/search?q=maria pag').get_json()
    assert [row['name'] for row in by_tokens] == ['Maria C Pagulayan']
    inactive = client.get('/people/search?q=Inactive').get_json()
    assert inactive == []
    assert client.get('/people/search?q=zz').get_json() == []

def test_contribution_csv_import(client):
    login(client, 'sally@example.org')
    template = client.get('/contributions/template.csv')
    assert template.status_code == 200
    assert b'date,first_name,last_name,email,fund,amount,method,fiscal_year,reference,note' in template.data
    csv_body = (
        'date,first_name,last_name,email,fund,amount,method,fiscal_year,reference,note\n'
        '2026-09-01,Ada,Member,,General Fund,40.00,Check,2026,chk-1,\n'
        '2026-09-02,Ada,Member,,Missions,15.50,Cash,2026,,\n'
        '2026-09-03,Ada,Member,,General Fund,not-a-number,Check,2026,,\n'
        '2026-09-01,Ada,Member,,General Fund,40.00,Check,2026,chk-1,\n'
    )
    response = client.post(
        '/contributions/import',
        data={'csrf_token': token(client), 'file': (io.BytesIO(csv_body.encode('utf-8')), 'gifts.csv')},
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'Contributions import report' in response.data
    assert b'Loaded successfully' in response.data
    assert b'>2<' in response.data or b'2</strong>' in response.data
    assert b'Failed to load' in response.data
    assert b'amount must be a positive number' in response.data
    assert b'duplicate of another row in this file' in response.data
    with client.application.app_context():
        gifts = db.session.scalars(db.select(Gift).order_by(Gift.id)).all()
        assert len(gifts) == 2
        assert float(gifts[0].amount) == 40.00
        assert gifts[1].fund.name == 'Missions Fund'
        assert gifts[1].method == 'Cash'
    failed = client.get('/import/report/failures.csv')
    assert failed.status_code == 200
    assert b'failure_reason' in failed.data

def test_people_csv_import_report(client):
    login(client, 'sally@example.org')
    csv_body = (
        'first_name,last_name,member_type,email,phone,address_line1,address_line2,city,state,zip_code\n'
        'Ada,Member,Member,ada@example.org,555-0100,1 Main St,,Green Brook,NJ,08812\n'
        ',Visitor,Non-member,ben@example.org,555-0101,2 Main St,,Plainfield,NJ,07060\n'
        'Ada,Member,Member,ada@example.org,555-0100,1 Main St,,Green Brook,NJ,08812\n'
        'Ben,Visitor,Non-member,ben@example.org,555-0101,2 Main St,,Plainfield,NJ,07060\n'
    )
    response = client.post(
        '/people/import',
        data={'csrf_token': token(client), 'file': (io.BytesIO(csv_body.encode('utf-8')), 'people.csv')},
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'Confirm people updates' in response.data
    assert b'Existing person records will be changed' in response.data
    assert b'ada@example.org' in response.data
    confirm = client.post(
        '/people/import/confirm',
        data={'csrf_token': token(client), 'action': 'confirm'},
        follow_redirects=True,
    )
    assert b'People import report' in confirm.data
    assert b'first name and last name are required' in confirm.data
    assert b'duplicate of another row in this file' in confirm.data
    with client.application.app_context():
        people = db.session.scalars(db.select(Person).order_by(Person.id)).all()
        # fixture Ada updated + Ben created
        assert len(people) == 2
        ada = next(person for person in people if person.first_name == 'Ada')
        ben = next(person for person in people if person.first_name == 'Ben')
        assert ada.email == 'ada@example.org'
        assert ada.phone == '555-0100'
        assert ada.address_line1 == '1 Main St'
        assert ben.email == 'ben@example.org'

def test_people_csv_import_merge_rules_and_cancel(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        ada = db.session.get(Person, 1)
        ada.email = 'keep@example.org'
        ada.phone = ''
        ada.address_line1 = 'Old Street'
        ada.city = 'Green Brook'
        ada.state = 'NJ'
        ada.zip_code = '08812'
        db.session.commit()
    csv_body = (
        'first_name,last_name,member_type,email,phone,address_line1,address_line2,city,state,zip_code\n'
        'Ada,Member,Member,,555-9999,New Street,,,\n'
    )
    preview = client.post(
        '/people/import',
        data={'csrf_token': token(client), 'file': (io.BytesIO(csv_body.encode('utf-8')), 'merge.csv')},
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'Confirm people updates' in preview.data
    assert b'Phone' in preview.data and b'555-9999' in preview.data
    assert b'Address line 1' in preview.data and b'New Street' in preview.data
    canceled = client.post(
        '/people/import/confirm',
        data={'csrf_token': token(client), 'action': 'cancel'},
        follow_redirects=True,
    )
    assert b'No person records were changed' in canceled.data
    with client.application.app_context():
        ada = db.session.get(Person, 1)
        assert ada.email == 'keep@example.org'
        assert ada.phone == ''
        assert ada.address_line1 == 'Old Street'
    preview = client.post(
        '/people/import',
        data={'csrf_token': token(client), 'file': (io.BytesIO(csv_body.encode('utf-8')), 'merge.csv')},
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'Confirm people updates' in preview.data
    confirmed = client.post(
        '/people/import/confirm',
        data={'csrf_token': token(client), 'action': 'confirm'},
        follow_redirects=True,
    )
    assert b'People import report' in confirmed.data
    with client.application.app_context():
        ada = db.session.get(Person, 1)
        # blank CSV email keeps portal email; blank city/state/zip keep portal; phone filled; street overwritten
        assert ada.email == 'keep@example.org'
        assert ada.phone == '555-9999'
        assert ada.address_line1 == 'New Street'
        assert ada.city == 'Green Brook'
        assert ada.state == 'NJ'
        assert ada.zip_code == '08812'

def test_empty_csv_import_report(client):
    login(client, 'sally@example.org')
    response = client.post(
        '/people/import',
        data={'csrf_token': token(client), 'file': (io.BytesIO(b''), 'empty.csv')},
        content_type='multipart/form-data',
        follow_redirects=True,
    )
    assert b'People import report' in response.data
    assert b'The file is empty.' in response.data

def test_people_sort_and_pagination(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        db.session.add_all([
            Person(first_name=f'Person{i:02d}', last_name=f'Zed{i:02d}' if i < 20 else f'Able{i:02d}', member_type='Member', email=f'p{i}@example.org')
            for i in range(30)
        ])
        db.session.commit()
    default = client.get('/people').data.decode()
    assert 'sort=name' in default and '▲' in default
    assert default.count('<tr>') == 26  # header + 25 rows
    assert 'aria-label="People pages"' in default
    assert 'Back' in default and 'Next' in default
    assert '1-25 of 31' in default
    assert 'pager-goto' in default and '>Go<' in default
    assert 'Result per page' in default
    assert 'name="per_page"' in default
    assert '>10<' in default and '>25<' in default and '>50<' in default and '>100<' in default and '>250<' in default
    page2 = client.get('/people?page=2').data.decode()
    assert page2.count('<tr>') == 7  # header + 6 remaining (31 total with fixture Ada)
    assert 'aria-current="page">2<' in page2
    assert '26-31 of 31' in page2
    per50 = client.get('/people?per_page=50').data.decode()
    assert per50.count('<tr>') == 32  # header + all 31
    assert '1-31 of 31' in per50
    by_contact = client.get('/people?sort=contact&dir=asc').data.decode()
    assert 'sort=contact' in by_contact

def test_contributions_pagination(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        gifts = [
            Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026, 1, 1 + (i % 28)), amount=10 + i, method='Cash', entered_by_id=1)
            for i in range(30)
        ]
        db.session.add_all(gifts)
        db.session.commit()
    page1 = client.get('/contributions').data.decode()
    assert 'aria-label="Contributions pages"' in page1
    assert 'Back' in page1 and 'Next' in page1
    assert '1-25 of 30' in page1
    assert 'Result per page' in page1
    assert page1.count('<tr>') == 26  # header + 25 rows
    assert 'page=2' in page1
    page2 = client.get('/contributions?page=2').data.decode()
    assert page2.count('<tr>') == 6  # header + 5 remaining
    assert 'aria-current="page">2<' in page2
    assert '26-30 of 30' in page2
    ten = client.get('/contributions?per_page=10').data.decode()
    assert ten.count('<tr>') == 11
    assert '1-10 of 30' in ten

def test_reports_pagination_keeps_filters(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        gifts = [
            Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026, 2, 1 + (i % 27)), amount=5 + i, method='Cash', entered_by_id=1)
            for i in range(40)
        ]
        db.session.add_all(gifts)
        db.session.commit()
    filtered = client.get('/reports?year_id=1&start_date=2026-02-01&end_date=2026-02-28').data.decode()
    assert 'aria-label="Reports pages"' in filtered
    assert 'Contribution detail' in filtered
    assert '1-25 of' in filtered
    assert 'Result per page' in filtered
    assert 'year_id=1' in filtered and 'start_date=2026-02-01' in filtered
    page2 = client.get('/reports?year_id=1&start_date=2026-02-01&end_date=2026-02-28&page=2').data.decode()
    assert 'aria-current="page">2<' in page2
    assert 'year_id=1' in page2 and 'start_date=2026-02-01' in page2
    assert 'pager-goto' in page2
    fifty = client.get('/reports?year_id=1&start_date=2026-02-01&end_date=2026-02-28&per_page=50').data.decode()
    assert '1-40 of 40' in fifty or '1-' in fifty and 'of 40' in fifty
    assert 'per_page=50' in fifty

def test_contributions_sortable_default_date_desc(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        db.session.add(Person(first_name='Ben', last_name='Visitor', member_type='Non-member'))
        db.session.commit()
        db.session.add_all([
            Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,8,1), amount=10, method='Cash', entered_by_id=1),
            Gift(person_id=2, fiscal_year_id=1, fund_id=1, date=date(2026,10,1), amount=50, method='Check', entered_by_id=1),
        ])
        db.session.commit()
    default = client.get('/contributions').data.decode()
    assert 'sort=date' in default and '▼' in default
    assert default.index('2026-10-01') < default.index('2026-08-01')
    by_name = client.get('/contributions?sort=contributor&dir=asc').data.decode()
    assert by_name.index('Ada Member') < by_name.index('Ben Visitor')
    by_amount = client.get('/contributions?sort=amount&dir=desc').data.decode()
    assert by_amount.index('$50.00') < by_amount.index('$10.00')

def test_utilities_boa_people_csv(client):
    login(client, 'sally@example.org')
    assert b'People List from a BoA CSV' in client.get('/utilities').data
    sample = (
        'Date,Description,Amount,Running Bal.\n'
        '01/02/2026,"Zelle payment from MARIA MARANA for ""Tithes""; Conf# lkadzmzj9","250.00","100.00"\n'
        '01/02/2026,"Zelle payment from MARIA C PAGULAYAN for ""Kris tithes""; Conf# xyz","50.00","150.00"\n'
        '01/02/2026,"Zelle payment from MA DARIEANA J LAUD for ""offering""; Conf# qwe","40.00","190.00"\n'
        '01/02/2026,"Zelle payment from KIRSTEIN JASPER BALAO for ""tithe""; Conf# abc","76.00","266.00"\n'
        '01/05/2026,"Zelle payment from ELISEO PAGULAYAN Conf# f46mbhm9r","300.00","566.00"\n'
        '01/05/2026,"Zelle Recurring payment from ARTHUR AGULTO for ""Weekly giving""; Conf# j521","264.00","830.00"\n'
        '01/09/2026,"Zelle payment to Jayvee Pasturan Conf# yrjovwgsd","-96.00","734.00"\n'
        '01/05/2026,"Deposit","590.00","1324.00"\n'
    )
    response = client.post(
        '/utilities/boa-people',
        data={'csrf_token': token(client), 'file': (io.BytesIO(sample.encode('utf-8')), 'boa.csv')},
        content_type='multipart/form-data',
    )
    assert response.status_code == 200
    assert response.headers['Content-Disposition'].startswith('attachment;')
    body = response.data.decode()
    assert 'first_name,last_name' in body
    assert 'MARIA,MARANA' in body
    assert 'MARIA C,PAGULAYAN' in body
    assert 'MA DARIEANA,LAUD' in body
    assert 'KIRSTEIN,BALAO' in body
    assert 'ELISEO,PAGULAYAN' in body
    assert 'ARTHUR,AGULTO' in body
    assert 'Jayvee' not in body
    assert 'Pasturan' not in body

def test_utilities_boa_giving_csv(client):
    login(client, 'sally@example.org')
    utilities = client.get('/utilities').data
    assert b'Giving List from BoA General Fund CSV' in utilities
    assert b'title="Search and open a CSV file downloaded from BoA website"' in utilities
    sample = (
        'Date,Description,Amount,Running Bal.\n'
        '01/02/2026,"Zelle payment from MARIA MARANA for ""Tithes""; Conf# lkadzmzj9","250.00","100.00"\n'
        '01/02/2026,"Zelle payment from MARIA C PAGULAYAN for ""Kris tithes""; Conf# xyz","1,050.50","1150.50"\n'
        '01/05/2026,"Zelle payment from ELISEO PAGULAYAN Conf# f46mbhm9r","300.00","1450.50"\n'
        '01/09/2026,"Zelle payment to Jayvee Pasturan Conf# yrjovwgsd","-96.00","1354.50"\n'
        '01/05/2026,"Deposit","590.00","1944.50"\n'
    )
    response = client.post(
        '/utilities/boa-giving',
        data={'csrf_token': token(client), 'file': (io.BytesIO(sample.encode('utf-8')), 'boa.csv')},
        content_type='multipart/form-data',
    )
    assert response.status_code == 200
    assert response.headers['Content-Disposition'].startswith('attachment;')
    body = response.data.decode().strip().splitlines()
    assert body[0] == 'date,first_name,last_name,email,fund,amount,method,fiscal_year,reference,note'
    assert '2026-01-02,MARIA,MARANA,,General Fund,250.00,Online,,,Tithes' in body
    assert '2026-01-02,MARIA C,PAGULAYAN,,General Fund,1050.50,Online,,,Kris tithes' in body
    assert '2026-01-05,ELISEO,PAGULAYAN,,General Fund,300.00,Online,,,' in body
    assert 'Jayvee' not in response.data.decode()
    assert 'Pasturan' not in response.data.decode()

def test_person_delete_from_list(client):
    login(client, 'sally@example.org')
    listing = client.get('/people').data
    assert b'Delete' in listing
    with client.application.app_context():
        db.session.add(Person(first_name='Ben', last_name='Visitor', member_type='Non-member'))
        db.session.add(Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,9,1), amount=10, method='Cash', entered_by_id=1))
        db.session.commit()
        person_id = db.session.scalar(db.select(Person.id).filter_by(last_name='Visitor'))
    blocked = client.post('/people/1/delete', data={'csrf_token': token(client), 'audit_reason': 'Cleanup'}, follow_redirects=True)
    assert b'Cannot delete a person who has contribution records' in blocked.data
    deleted = client.post(f'/people/{person_id}/delete', data={'csrf_token': token(client), 'audit_reason': 'Duplicate record'}, follow_redirects=True)
    assert b'Ben Visitor deleted.' in deleted.data
    with client.application.app_context():
        assert db.session.get(Person, person_id) is None

def test_gift_delete_from_list(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        db.session.add(Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,9,1), amount=25, method='Check', entered_by_id=1))
        db.session.commit()
    listing = client.get('/contributions').data
    assert b'Delete' in listing
    deleted = client.post('/contributions/1/delete', data={'csrf_token': token(client), 'audit_reason': 'Entered twice'}, follow_redirects=True)
    assert b'Contribution deleted.' in deleted.data
    with client.application.app_context():
        assert db.session.get(Gift, 1) is None

def test_reconciliation_is_audited_and_role_limited(client):
    login(client, 'sally@example.org')
    with client.application.app_context():
        db.session.add(Gift(person_id=1, fiscal_year_id=1, fund_id=1, date=date(2026,9,1), amount=25, method='Check', entered_by_id=1)); db.session.commit()
    response = client.post('/contributions/1/reconcile', data={'csrf_token':token(client), 'audit_reason': 'Matched bank deposit'}, follow_redirects=True)
    assert b'Reconciliation status updated' in response.data
    with client.application.app_context(): assert db.session.get(Gift,1).reconciled is True
    audit_page = client.get('/admin/audit').data
    assert b'Matched bank deposit' in audit_page
    assert 'Reconciled: No → Yes'.encode() in audit_page

"""GBBC Finance and Giving Portal. Configure DATABASE_URL for MySQL in production."""
import csv
import hashlib
import io
import json
import os
import re
import secrets
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from functools import wraps
from pathlib import Path
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

import click
from flask import Flask, abort, current_app, flash, g, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from mailmerge import MailMerge
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader, simpleSplit
from reportlab.pdfgen import canvas
from werkzeug.security import check_password_hash, generate_password_hash

ROLES = ('Super Admin', 'System Administrator', 'Finance Staff', 'Church Leader', 'People-Manager', 'People Viewer')
SIGNUP_ROLES = ('System Administrator', 'Finance Staff', 'People-Manager', 'People Viewer')
LEGACY_ROLE_MAP = {
    'Admin': 'System Administrator',
    'Finance': 'Finance Staff',
}
INACTIVITY_TIMEOUT = timedelta(minutes=10)
INACTIVITY_TIMEOUT_SECONDS = int(INACTIVITY_TIMEOUT.total_seconds())
US_STATES = (
    ('AL', 'Alabama'), ('AK', 'Alaska'), ('AZ', 'Arizona'), ('AR', 'Arkansas'), ('CA', 'California'),
    ('CO', 'Colorado'), ('CT', 'Connecticut'), ('DE', 'Delaware'), ('DC', 'District of Columbia'),
    ('FL', 'Florida'), ('GA', 'Georgia'), ('HI', 'Hawaii'), ('ID', 'Idaho'), ('IL', 'Illinois'),
    ('IN', 'Indiana'), ('IA', 'Iowa'), ('KS', 'Kansas'), ('KY', 'Kentucky'), ('LA', 'Louisiana'),
    ('ME', 'Maine'), ('MD', 'Maryland'), ('MA', 'Massachusetts'), ('MI', 'Michigan'), ('MN', 'Minnesota'),
    ('MS', 'Mississippi'), ('MO', 'Missouri'), ('MT', 'Montana'), ('NE', 'Nebraska'), ('NV', 'Nevada'),
    ('NH', 'New Hampshire'), ('NJ', 'New Jersey'), ('NM', 'New Mexico'), ('NY', 'New York'),
    ('NC', 'North Carolina'), ('ND', 'North Dakota'), ('OH', 'Ohio'), ('OK', 'Oklahoma'), ('OR', 'Oregon'),
    ('PA', 'Pennsylvania'), ('RI', 'Rhode Island'), ('SC', 'South Carolina'), ('SD', 'South Dakota'),
    ('TN', 'Tennessee'), ('TX', 'Texas'), ('UT', 'Utah'), ('VT', 'Vermont'), ('VA', 'Virginia'),
    ('WA', 'Washington'), ('WV', 'West Virginia'), ('WI', 'Wisconsin'), ('WY', 'Wyoming'),
)
US_STATE_CODES = {code for code, _ in US_STATES}
DEFAULT_FUNDS = ('General Fund', 'Missions Fund', 'Building Fund', 'Benevolence Fund')
FUND_ALIASES = {
    'Missions': 'Missions Fund',
    'Benevolence': 'Benevolence Fund',
}
PERMISSIONS = {
    'Super Admin': {'overview', 'people_view', 'people_write', 'gift_view', 'gift_write', 'reports', 'detail_reports', 'statements', 'users', 'roles', 'fiscal', 'feedback', 'settings', 'feedback_review', 'audit_view', 'utilities', 'backup'},
    'System Administrator': {'overview', 'people_view', 'reports', 'users', 'roles', 'fiscal', 'feedback', 'settings', 'feedback_review', 'audit_view', 'utilities', 'backup'},
    'Finance Staff': {'overview', 'people_view', 'people_write', 'gift_view', 'gift_write', 'reports', 'detail_reports', 'statements', 'feedback', 'utilities'},
    'Church Leader': {'overview', 'people_view', 'reports', 'feedback', 'utilities'},
    'People-Manager': {'overview', 'people_view', 'people_write', 'feedback'},
    'People Viewer': {'people_view'},
}
PAGE_MATRIX = [
    ('Dashboard', 'overview'), ('People directory', 'people_view'), ('Create and edit people', 'people_write'),
    ('Contribution detail', 'gift_view'), ('Record and correct gifts', 'gift_write'), ('Aggregate giving reports', 'reports'),
    ('Individual giving reports', 'detail_reports'), ('Giving statements', 'statements'),
    ('Utilities', 'utilities'),
    ('User administration and signup approvals', 'users'), ('Role permissions', 'roles'), ('Fiscal-year administration', 'fiscal'),
    ('Feedback', 'feedback'), ('Website settings', 'settings'),
    ('Review submitted feedback', 'feedback_review'), ('Audit history', 'audit_view'),
    ('Backup and restore', 'backup'),
]
BACKUP_SHEETS = ('People', 'Contributions', 'Users', 'Audit log')
BACKUP_MARK = 'GBBC_BACKUP_V1'

db = SQLAlchemy()

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    first_name = db.Column(db.String(100), nullable=False)
    last_name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(255), unique=True, nullable=False)
    username = db.Column(db.String(100), unique=True)
    phone = db.Column(db.String(50), nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(40), nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)
    last_login = db.Column(db.DateTime)
    @property
    def name(self): return f'{self.first_name} {self.last_name}'

class Person(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    first_name = db.Column(db.String(100), nullable=False)
    last_name = db.Column(db.String(100), nullable=False)
    member_type = db.Column(db.String(20), nullable=False)
    address_line1 = db.Column(db.String(120), default='')
    address_line2 = db.Column(db.String(120), default='')
    city = db.Column(db.String(100), default='')
    state = db.Column(db.String(2), default='')
    zip_code = db.Column(db.String(20), default='')
    email = db.Column(db.String(255), default='')
    phone = db.Column(db.String(50), default='')
    active = db.Column(db.Boolean, default=True, nullable=False)
    @property
    def name(self): return f'{self.first_name} {self.last_name}'
    @property
    def address_display(self):
        return ', '.join(person_address_lines(self)) or ''

class FiscalYear(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(30), unique=True, nullable=False)
    start_date = db.Column(db.Date, nullable=False)
    end_date = db.Column(db.Date, nullable=False)
    is_open = db.Column(db.Boolean, default=True, nullable=False)

class Fund(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), unique=True, nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)

class Gift(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    person_id = db.Column(db.Integer, db.ForeignKey('person.id'), nullable=False)
    person = db.relationship('Person')
    fiscal_year_id = db.Column(db.Integer, db.ForeignKey('fiscal_year.id'), nullable=False)
    fiscal_year = db.relationship('FiscalYear')
    fund_id = db.Column(db.Integer, db.ForeignKey('fund.id'), nullable=False)
    fund = db.relationship('Fund')
    date = db.Column(db.Date, nullable=False)
    amount = db.Column(db.Numeric(12, 2), nullable=False)
    method = db.Column(db.String(40), nullable=False)
    reference = db.Column(db.String(120), default='')
    note = db.Column(db.Text, default='')
    status = db.Column(db.String(20), default='active', nullable=False)
    reconciled = db.Column(db.Boolean, default=False, nullable=False)
    entered_by_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    entered_by = db.relationship('User')
    created_at = db.Column(db.DateTime, default=lambda: utcnow(), nullable=False)

class Feedback(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    user = db.relationship('User')
    category = db.Column(db.String(30), nullable=False)
    subject = db.Column(db.String(150), nullable=False)
    message = db.Column(db.Text, nullable=False)
    status = db.Column(db.String(30), default='new')
    created_at = db.Column(db.DateTime, default=lambda: utcnow(), nullable=False)

class Audit(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'))
    action = db.Column(db.String(100), nullable=False)
    entity = db.Column(db.String(60), nullable=False)
    entity_id = db.Column(db.Integer)
    timestamp = db.Column(db.DateTime, default=lambda: utcnow(), nullable=False)
    summary = db.Column(db.String(255), default='')
    reason = db.Column(db.String(255), default='')
    details = db.Column(db.Text, default='')

class Setting(db.Model):
    key = db.Column(db.String(80), primary_key=True)
    value = db.Column(db.Text, nullable=False)

class LoginAttempt(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(255), nullable=False, index=True)
    attempted_at = db.Column(db.DateTime, default=lambda: utcnow(), nullable=False)

class SignupRequest(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    first_name = db.Column(db.String(32), nullable=False)
    last_name = db.Column(db.String(32), nullable=False)
    email = db.Column(db.String(255), nullable=False, index=True)
    password_hash = db.Column(db.String(255), nullable=False)
    requested_role = db.Column(db.String(40), nullable=False)
    status = db.Column(db.String(20), default='pending', nullable=False, index=True)
    created_at = db.Column(db.DateTime, default=lambda: utcnow(), nullable=False)
    reviewed_at = db.Column(db.DateTime)
    reviewed_by_id = db.Column(db.Integer, db.ForeignKey('user.id'))
    reviewed_by = db.relationship('User')
    decision_note = db.Column(db.String(255), default='')
    @property
    def name(self):
        return f'{self.first_name} {self.last_name}'

def utcnow():
    return datetime.now(timezone.utc).replace(tzinfo=None)

def hash_password(password):
    method = 'scrypt' if hasattr(hashlib, 'scrypt') else 'pbkdf2:sha256'
    return generate_password_hash(password, method=method)

def validate_signup_password(password):
    if len(password) < 8:
        raise ValueError('Password must be at least 8 characters.')
    if not re.search(r'[A-Z]', password):
        raise ValueError('Password must include at least one capital letter.')
    if not re.search(r'[a-z]', password):
        raise ValueError('Password must include at least one lowercase letter.')
    if not re.search(r'[0-9]', password):
        raise ValueError('Password must include at least one number.')
    if not re.search(r'[^A-Za-z0-9]', password):
        raise ValueError('Password must include at least one special character.')

def is_local_path(target):
    # Browsers treat backslashes as slashes, so "/\evil.com" is an off-site link.
    if not target.startswith('/') or target[1:2] in ('/', '\\'):
        return False
    parsed = urlparse(target.replace('\\', '/'))
    return not parsed.scheme and not parsed.netloc and not any(ord(ch) < 32 for ch in target)

def default_home_path(user):
    perms = PERMISSIONS.get(user.role, set())
    if 'overview' in perms:
        return '/'
    if 'people_view' in perms:
        return '/people'
    return '/help'

MAIL_SETTING_KEYS = (
    'mail_from', 'smtp_host', 'smtp_port', 'smtp_username', 'smtp_password', 'smtp_encryption',
)

def get_setting(key, default=''):
    row = db.session.get(Setting, key)
    return row.value if row and row.value is not None else default

def set_setting(key, value):
    row = db.session.get(Setting, key)
    if row is None:
        db.session.add(Setting(key=key, value=value))
    else:
        row.value = value

def portal_mail_config():
    """SMTP settings: environment variables override Website settings values."""
    def env_or_setting(env_name, setting_key, default=''):
        return os.environ.get(env_name, '').strip() or get_setting(setting_key, default).strip()

    encryption = env_or_setting('SMTP_ENCRYPTION', 'smtp_encryption', '')
    if not encryption:
        if os.environ.get('SMTP_SSL', '').strip() == '1':
            encryption = 'ssl'
        elif os.environ.get('SMTP_STARTTLS', '1').strip() == '0':
            encryption = 'none'
        else:
            encryption = 'starttls'
    return {
        'from_addr': env_or_setting('MAIL_FROM', 'mail_from', 'noreply@greenbrookchurch.org'),
        'host': env_or_setting('SMTP_HOST', 'smtp_host'),
        'port': int(env_or_setting('SMTP_PORT', 'smtp_port', '587') or '587'),
        'username': env_or_setting('SMTP_USERNAME', 'smtp_username'),
        'password': env_or_setting('SMTP_PASSWORD', 'smtp_password'),
        'encryption': encryption.lower(),
    }

def portal_from_address():
    return portal_mail_config()['from_addr'] or 'noreply@greenbrookchurch.org'

def mail_is_configured():
    return bool(portal_mail_config()['host'])

def _write_mail_outbox(message, outbox_name):
    outbox = Path(current_app.instance_path) / 'mail_outbox'
    outbox.mkdir(parents=True, exist_ok=True)
    stamp = utcnow().strftime('%Y%m%d%H%M%S')
    safe_name = re.sub(r'[^A-Za-z0-9_.-]+', '_', outbox_name)[:80]
    path = outbox / f'{safe_name}_{stamp}.txt'
    path.write_text(message, encoding='utf-8')
    return path

def _smtp_help_message(exc, config):
    text = str(exc).strip() or exc.__class__.__name__
    host = (config.get('host') or '').lower()
    tips = []
    if 'gmail.com' in host or 'googlemail.com' in host:
        tips.append(
            'For Gmail, use an App Password (Google Account → Security → 2-Step Verification → App passwords), not your normal login password.'
        )
        tips.append('Recommended settings: smtp.gmail.com, port 587, STARTTLS, and From address = your Gmail address.')
        if config.get('from_addr') and config.get('username'):
            from_l = config['from_addr'].lower()
            user_l = config['username'].lower()
            if from_l != user_l and not from_l.endswith('@gmail.com') and not from_l.endswith('@googlemail.com'):
                tips.append(
                    f'Gmail usually rejects From addresses that are not the signed-in account. '
                    f'Set From to {config["username"]} (or send as that Gmail address).'
                )
    if 'connection unexpectedly closed' in text.lower():
        tips.append('Try port 587 with STARTTLS if SSL on 465 keeps dropping.')
    if tips:
        return f'{text} — ' + ' '.join(tips)
    return text

def send_portal_email(to_addr, subject, body, from_addr=None, outbox_name='mail'):
    """Send email via SMTP when configured; otherwise write to instance/mail_outbox."""
    import smtplib
    from email.message import EmailMessage

    config = portal_mail_config()
    from_addr = (from_addr or config['from_addr'] or 'noreply@greenbrookchurch.org').strip()
    # Gmail requires the From address to match the authenticated mailbox unless "Send mail as" is set up.
    host = (config['host'] or '').lower()
    username = (config['username'] or '').strip()
    if username and ('gmail.com' in host or 'googlemail.com' in host):
        user_l = username.lower()
        from_l = from_addr.lower()
        if from_l != user_l and not from_l.endswith('@gmail.com') and not from_l.endswith('@googlemail.com'):
            from_addr = username
    message = f'From: {from_addr}\nTo: {to_addr}\nSubject: {subject}\n\n{body}'
    if not config['host']:
        _write_mail_outbox(message, outbox_name)
        return 'queued'

    email_msg = EmailMessage()
    email_msg['From'] = from_addr
    email_msg['To'] = to_addr
    email_msg['Subject'] = subject
    if config['from_addr'] and config['from_addr'].strip().lower() != from_addr.lower():
        email_msg['Reply-To'] = config['from_addr'].strip()
    email_msg.set_content(body)
    try:
        if config['encryption'] == 'ssl':
            context = smtplib.SMTP_SSL(config['host'], config['port'], timeout=30)
        else:
            context = smtplib.SMTP(config['host'], config['port'], timeout=30)
        with context as smtp:
            smtp.ehlo()
            if config['encryption'] == 'starttls':
                smtp.starttls()
                smtp.ehlo()
            if username:
                if not config['password']:
                    raise RuntimeError('SMTP password is missing. Save an App Password in Website settings.')
                smtp.login(username, config['password'])
            smtp.send_message(email_msg)
        return 'sent'
    except Exception as exc:
        _write_mail_outbox(message, f'{outbox_name}_failed')
        raise RuntimeError(_smtp_help_message(exc, {**config, 'from_addr': from_addr})) from exc

def notify_signup_received(signup):
    """Acknowledge that a signup request was received and is awaiting approval."""
    subject = 'GBBC Finance signup request received'
    body = (
        f'Dear {signup.first_name},\n\n'
        f'We received your GBBC Finance portal signup request for the role '
        f'"{signup.requested_role}".\n\n'
        f'This email confirms we have your request. An administrator must approve '
        f'it before you can sign in. You will receive another email with the '
        f'approval or decline decision.\n\n'
        f'Green Brook Baptist Church\n'
    )
    return send_portal_email(
        signup.email,
        subject,
        body,
        outbox_name=f'signup_{signup.id}_received',
    )

def notify_signup_decision(signup, approved, reviewer):
    """Send approval/decline notice from the reviewing admin's email (or local outbox fallback)."""
    decision = 'approved' if approved else 'declined'
    subject = f'GBBC Finance account request {decision}'
    if approved:
        body = (
            f'Dear {signup.first_name},\n\n'
            f'Your GBBC Finance portal account request has been approved by {reviewer.name}.\n'
            f'Role: {signup.requested_role}\n'
            f'Sign in with your email address: {signup.email}\n\n'
            f'Green Brook Baptist Church\n'
        )
    else:
        body = (
            f'Dear {signup.first_name},\n\n'
            f'Your GBBC Finance portal account request has been declined by {reviewer.name}.\n'
            f'If you believe this was a mistake, contact your church administrator.\n\n'
            f'Green Brook Baptist Church\n'
        )
    return send_portal_email(
        signup.email,
        subject,
        body,
        from_addr=reviewer.email,
        outbox_name=f'signup_{signup.id}_{decision}',
    )

def excel_safe(value):
    text = str(value)
    return "'" + text if text.startswith(('=', '+', '-', '@', '\t', '\r')) else text

EASTERN_TZ = ZoneInfo('America/New_York')

def format_eastern(dt):
    """Format a naive-UTC or aware datetime in America/New_York."""
    if dt is None:
        return ''
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(EASTERN_TZ).strftime('%Y-%m-%d %H:%M %Z')

def audit_display_value(value):
    if value is None:
        return '(empty)'
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, Decimal):
        return f'{value:.2f}'
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.isoformat()
    text = str(value).strip()
    return text if text else '(empty)'

def format_audit_changes(changes):
    """changes: iterable of (label, old_value, new_value)."""
    lines = []
    for label, old, new in changes:
        old_text = audit_display_value(old)
        new_text = audit_display_value(new)
        if old_text == new_text:
            continue
        lines.append(f'{label}: {old_text} → {new_text}')
    return '\n'.join(lines)

def require_audit_reason(required=True):
    reason = (request.form.get('audit_reason') or '').strip()
    if required and not reason:
        raise ValueError('A reason for this change is required.')
    return reason[:255]

def audit(action, entity, entity_id=None, summary='', reason='', details='', changes=None):
    if changes is not None and not details:
        details = format_audit_changes(changes)
    db.session.add(Audit(
        user_id=g.user.id if g.user else None,
        action=action,
        entity=entity,
        entity_id=entity_id,
        summary=(summary or '')[:255],
        reason=(reason or '')[:255],
        details=details or '',
    ))

def _sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(cell or '').strip() for cell in rows[0]]
    data = []
    for raw in rows[1:]:
        if raw is None or all(cell is None or str(cell).strip() == '' for cell in raw):
            continue
        data.append({headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))})
    return headers, data

def _cell_str(value):
    if value is None:
        return ''
    return str(value).strip()

def _cell_bool(value, default=False):
    if value is None or value == '':
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in ('1', 'true', 'yes', 'y', 'active'):
        return True
    if text in ('0', 'false', 'no', 'n', 'inactive'):
        return False
    return default

def _cell_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])

def _cell_datetime(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    text = str(value).strip().replace('Z', '')
    if 'T' in text:
        return datetime.fromisoformat(text)
    if ' ' in text:
        return datetime.fromisoformat(text)
    return datetime.fromisoformat(f'{text}T00:00:00')

def _cell_amount(value):
    """Parse money from Excel/backup cells without float-precision false failures."""
    if value is None or value == '':
        raise ValueError('Enter a positive amount with no more than two decimal places.')
    try:
        if isinstance(value, Decimal):
            amount = value
        elif isinstance(value, bool):
            raise ValueError
        elif isinstance(value, int):
            amount = Decimal(value)
        elif isinstance(value, float):
            # Excel often returns floats; round via string to avoid binary artifacts.
            amount = Decimal(f'{value:.2f}')
        else:
            text = str(value).strip().replace(',', '').replace('$', '')
            amount = Decimal(text)
        if not amount.is_finite() or amount <= 0:
            raise ValueError
        return amount.quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError('Enter a positive amount with no more than two decimal places.')

def fiscal_year_audit_window(year):
    start = datetime.combine(year.start_date, datetime.min.time())
    end = datetime.combine(year.end_date, datetime.max.time()).replace(microsecond=0)
    return start, end

def build_year_backup_workbook(year):
    """Build a 4-sheet Excel backup for one fiscal year."""
    gifts = db.session.scalars(
        db.select(Gift).where(Gift.fiscal_year_id == year.id).order_by(Gift.date, Gift.id)
    ).all()
    people = db.session.scalars(db.select(Person).order_by(Person.id)).all()
    users = db.session.scalars(db.select(User).order_by(User.id)).all()
    audit_start, audit_end = fiscal_year_audit_window(year)
    events = db.session.scalars(
        db.select(Audit).where(Audit.timestamp >= audit_start, Audit.timestamp <= audit_end).order_by(Audit.id)
    ).all()

    wb = Workbook()
    wb.properties.title = f'GBBC Finance backup {year.name}'
    wb.properties.subject = f'{BACKUP_MARK}|{year.name}|{year.id}|{year.start_date.isoformat()}|{year.end_date.isoformat()}'

    ws_people = wb.active
    ws_people.title = 'People'
    ws_people.append([
        'id', 'first_name', 'last_name', 'member_type', 'address_line1', 'address_line2',
        'city', 'state', 'zip_code', 'email', 'phone', 'active',
    ])
    for person in people:
        ws_people.append([
            person.id, person.first_name, person.last_name, person.member_type,
            person.address_line1 or '', person.address_line2 or '', person.city or '',
            person.state or '', person.zip_code or '', person.email or '', person.phone or '',
            'yes' if person.active else 'no',
        ])

    ws_gifts = wb.create_sheet('Contributions')
    ws_gifts.append([
        'id', 'person_id', 'person_name', 'fiscal_year', 'fund', 'date', 'amount', 'method',
        'reference', 'note', 'status', 'reconciled', 'entered_by_email', 'created_at',
    ])
    for gift in gifts:
        ws_gifts.append([
            gift.id, gift.person_id, gift.person.name, year.name, gift.fund.name,
            gift.date.isoformat(), f'{gift.amount:.2f}', gift.method, gift.reference or '',
            gift.note or '', gift.status, 'yes' if gift.reconciled else 'no',
            gift.entered_by.email if gift.entered_by else '',
            gift.created_at.isoformat(sep=' ', timespec='seconds') if gift.created_at else '',
        ])

    ws_users = wb.create_sheet('Users')
    ws_users.append([
        'id', 'first_name', 'last_name', 'email', 'username', 'phone', 'role', 'active',
        'password_hash', 'last_login',
    ])
    for user in users:
        ws_users.append([
            user.id, user.first_name, user.last_name, user.email, user.username or '',
            user.phone or '', user.role, 'yes' if user.active else 'no', user.password_hash,
            user.last_login.isoformat(sep=' ', timespec='seconds') if user.last_login else '',
        ])

    ws_audit = wb.create_sheet('Audit log')
    ws_audit.append([
        'id', 'user_email', 'action', 'entity', 'entity_id', 'timestamp', 'summary', 'reason', 'details',
    ])
    for event in events:
        actor = db.session.get(User, event.user_id) if event.user_id else None
        ws_audit.append([
            event.id,
            actor.email if actor else '',
            event.action,
            event.entity,
            event.entity_id if event.entity_id is not None else '',
            event.timestamp.isoformat(sep=' ', timespec='seconds') if event.timestamp else '',
            event.summary or '',
            event.reason or '',
            event.details or '',
        ])

    return wb, {
        'people': len(people),
        'contributions': len(gifts),
        'users': len(users),
        'audit': len(events),
        'year_name': year.name,
        'year_id': year.id,
    }

def parse_backup_workbook(file_storage_or_path, expected_year):
    """Validate a backup workbook and return sheet row dicts plus summary."""
    wb = load_workbook(file_storage_or_path, data_only=True)
    missing = [name for name in BACKUP_SHEETS if name not in wb.sheetnames]
    if missing:
        raise ValueError(f'Backup file is missing worksheet(s): {", ".join(missing)}')
    subject = (wb.properties.subject or '').strip()
    if subject.startswith(BACKUP_MARK + '|'):
        parts = subject.split('|')
        if len(parts) >= 3 and parts[1] != expected_year.name:
            raise ValueError(
                f'This backup is for fiscal year {parts[1]}, but {expected_year.name} was selected.'
            )
    people_headers, people_rows = _sheet_rows(wb['People'])
    gift_headers, gift_rows = _sheet_rows(wb['Contributions'])
    user_headers, user_rows = _sheet_rows(wb['Users'])
    audit_headers, audit_rows = _sheet_rows(wb['Audit log'])
    required = {
        'People': {'id', 'first_name', 'last_name', 'member_type'},
        'Contributions': {'person_id', 'fund', 'date', 'amount', 'method', 'fiscal_year'},
        'Users': {'id', 'first_name', 'last_name', 'email', 'role', 'password_hash'},
        'Audit log': {'action', 'entity', 'timestamp'},
    }
    for sheet_name, headers in (
        ('People', people_headers),
        ('Contributions', gift_headers),
        ('Users', user_headers),
        ('Audit log', audit_headers),
    ):
        missing_cols = required[sheet_name] - set(headers)
        if missing_cols:
            raise ValueError(f'{sheet_name} sheet is missing column(s): {", ".join(sorted(missing_cols))}')
    for row in gift_rows:
        fy = _cell_str(row.get('fiscal_year'))
        if fy and fy != expected_year.name:
            raise ValueError(
                f'Contribution rows reference fiscal year {fy}, but {expected_year.name} was selected.'
            )
    return {
        'people': people_rows,
        'contributions': gift_rows,
        'users': user_rows,
        'audit': audit_rows,
        'summary': {
            'people': len(people_rows),
            'contributions': len(gift_rows),
            'users': len(user_rows),
            'audit': len(audit_rows),
            'year_name': expected_year.name,
            'year_id': expected_year.id,
        },
    }

def apply_year_backup_restore(payload, year, actor):
    """Restore people/users upserts and replace year contributions + year audit rows."""
    # People upsert
    for row in payload['people']:
        person_id = int(row['id']) if row.get('id') not in (None, '') else None
        person = db.session.get(Person, person_id) if person_id else None
        if person is None:
            person = Person(id=person_id) if person_id else Person()
            db.session.add(person)
        person.first_name = _cell_str(row.get('first_name'))[:100]
        person.last_name = _cell_str(row.get('last_name'))[:100]
        person.member_type = _cell_str(row.get('member_type')) or 'Non-member'
        if person.member_type not in ('Member', 'Non-member'):
            person.member_type = 'Non-member'
        person.address_line1 = _cell_str(row.get('address_line1'))[:120]
        person.address_line2 = _cell_str(row.get('address_line2'))[:120]
        person.city = _cell_str(row.get('city'))[:100]
        person.state = _cell_str(row.get('state'))[:2].upper()
        person.zip_code = _cell_str(row.get('zip_code'))[:20]
        person.email = _cell_str(row.get('email'))[:255]
        person.phone = _cell_str(row.get('phone'))[:50]
        person.active = _cell_bool(row.get('active'), True)
    db.session.flush()

    # Users upsert (keep the acting administrator signed in and active)
    for row in payload['users']:
        email = _cell_str(row.get('email')).lower()
        if not email or not valid_email(email):
            continue
        user_id = int(row['id']) if row.get('id') not in (None, '') else None
        user = db.session.get(User, user_id) if user_id else None
        if user is None:
            user = db.session.scalar(db.select(User).filter_by(email=email))
        if user is None:
            user = User(id=user_id) if user_id else User()
            db.session.add(user)
        if user.id == actor.id:
            # Never lock out the administrator performing the restore.
            user.first_name = _cell_str(row.get('first_name'))[:100] or user.first_name
            user.last_name = _cell_str(row.get('last_name'))[:100] or user.last_name
            user.phone = _cell_str(row.get('phone'))[:50] or user.phone
            user.active = True
            continue
        user.first_name = _cell_str(row.get('first_name'))[:100]
        user.last_name = _cell_str(row.get('last_name'))[:100]
        user.email = email
        username = _cell_str(row.get('username')) or None
        user.username = username.lower() if username else None
        user.phone = _cell_str(row.get('phone'))[:50]
        role = _cell_str(row.get('role'))
        user.role = role if role in ROLES else 'People Viewer'
        user.active = _cell_bool(row.get('active'), True)
        password_hash = _cell_str(row.get('password_hash'))
        if password_hash:
            user.password_hash = password_hash
        last_login = _cell_datetime(row.get('last_login'))
        if last_login:
            user.last_login = last_login
    db.session.flush()

    # Replace contributions for this fiscal year
    existing_gifts = db.session.scalars(db.select(Gift).where(Gift.fiscal_year_id == year.id)).all()
    for gift in existing_gifts:
        db.session.delete(gift)
    db.session.flush()
    for row in payload['contributions']:
        person_id = int(row['person_id'])
        person = db.session.get(Person, person_id)
        if not person:
            raise ValueError(f'Contribution references missing person_id {person_id}.')
        fund = resolve_fund_name(_cell_str(row.get('fund')))
        entered_email = _cell_str(row.get('entered_by_email')).lower()
        entered_by = db.session.scalar(db.select(User).filter_by(email=entered_email)) if entered_email else None
        if entered_by is None:
            entered_by = actor
        amount = _cell_amount(row.get('amount'))
        gift = Gift(
            person_id=person.id,
            fiscal_year_id=year.id,
            fund_id=fund.id,
            date=_cell_date(row.get('date')),
            amount=amount,
            method=_cell_str(row.get('method')) or 'Other',
            reference=_cell_str(row.get('reference'))[:120],
            note=_cell_str(row.get('note')),
            status=_cell_str(row.get('status')) or 'active',
            reconciled=_cell_bool(row.get('reconciled'), False),
            entered_by_id=entered_by.id,
        )
        if gift.status not in ('active', 'void'):
            gift.status = 'active'
        if gift.method not in ('Check', 'Cash', 'Online', 'Other'):
            gift.method = 'Other'
        created_at = _cell_datetime(row.get('created_at'))
        if created_at:
            gift.created_at = created_at
        db.session.add(gift)
    db.session.flush()

    # Replace audit rows that fall in the fiscal-year window
    audit_start, audit_end = fiscal_year_audit_window(year)
    old_events = db.session.scalars(
        db.select(Audit).where(Audit.timestamp >= audit_start, Audit.timestamp <= audit_end)
    ).all()
    for event in old_events:
        db.session.delete(event)
    db.session.flush()
    for row in payload['audit']:
        email = _cell_str(row.get('user_email')).lower()
        user = db.session.scalar(db.select(User).filter_by(email=email)) if email else None
        stamp = _cell_datetime(row.get('timestamp')) or utcnow()
        entity_id_raw = row.get('entity_id')
        entity_id = int(entity_id_raw) if entity_id_raw not in (None, '') else None
        db.session.add(Audit(
            user_id=user.id if user else None,
            action=_cell_str(row.get('action'))[:100] or 'restore',
            entity=_cell_str(row.get('entity'))[:60] or 'unknown',
            entity_id=entity_id,
            timestamp=stamp,
            summary=_cell_str(row.get('summary'))[:255],
            reason=_cell_str(row.get('reason'))[:255],
            details=_cell_str(row.get('details')),
        ))

def backup_pending_path(user_id):
    folder = Path(current_app.instance_path) / 'backup_pending'
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f'restore_{user_id}.xlsx'

def valid_email(value):
    return bool(re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', value))

def parse_amount(raw):
    try:
        value = Decimal(raw)
        if not value.is_finite() or value <= 0 or value.as_tuple().exponent < -2: raise ValueError
        return value.quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        raise ValueError('Enter a positive amount with no more than two decimal places.')

def person_address_lines(person):
    lines = []
    line1 = (person.address_line1 or '').strip()
    line2 = (person.address_line2 or '').strip()
    if line1:
        lines.append(line1)
    if line2:
        lines.append(line2)
    city = (person.city or '').strip()
    state = (person.state or '').strip()
    zip_code = (person.zip_code or '').strip()
    if city or state or zip_code:
        if city and state:
            locality = f'{city}, {state}'
        else:
            locality = city or state
        if zip_code:
            locality = f'{locality} {zip_code}'.strip()
        lines.append(locality)
    return lines

def normalize_state(value):
    code = (value or '').strip().upper()
    if not code:
        return ''
    if code in US_STATE_CODES:
        return code
    for state_code, state_name in US_STATES:
        if state_name.upper() == code:
            return state_code
    if re.fullmatch(r'[A-Z]{2}', code):
        return code
    raise ValueError('Enter a valid 2-letter state code.')

def normalize_zip(value):
    zip_code = (value or '').strip()
    if not zip_code:
        return ''
    if not re.fullmatch(r'\d{5}(-\d{4})?', zip_code):
        raise ValueError('Enter a ZIP code as 12345 or 12345-6789.')
    return zip_code

def read_person_address(source):
    """Read address fields from a form or CSV row dict."""
    get = source.get if hasattr(source, 'get') else lambda key, default='': source.get(key, default)
    line1 = (get('address_line1') or get('address') or '').strip()
    line2 = (get('address_line2') or '').strip()
    city = (get('city') or '').strip()
    state = normalize_state(get('state') or '')
    zip_code = normalize_zip(get('zip_code') or get('zip') or '')
    return line1, line2, city, state, zip_code

def migrate_legacy_roles():
    """Map retired Admin/Finance roles onto System Administrator/Finance Staff."""
    from sqlalchemy import inspect, text
    tables = set(inspect(db.engine).get_table_names())
    if 'user' in tables:
        for old_role, new_role in LEGACY_ROLE_MAP.items():
            db.session.execute(
                text('UPDATE user SET role = :new_role WHERE role = :old_role'),
                {'new_role': new_role, 'old_role': old_role},
            )
    if 'signup_request' in tables:
        for old_role, new_role in LEGACY_ROLE_MAP.items():
            db.session.execute(
                text('UPDATE signup_request SET requested_role = :new_role WHERE requested_role = :old_role'),
                {'new_role': new_role, 'old_role': old_role},
            )
    db.session.commit()

def ensure_schema():
    """Add columns introduced after the first release; migrate legacy person.address."""
    from sqlalchemy import inspect, text
    db.create_all()
    migrate_legacy_roles()
    inspector = inspect(db.engine)
    tables = set(inspector.get_table_names())
    if 'user' in tables:
        user_cols = {col['name'] for col in inspector.get_columns('user')}
        if 'username' not in user_cols:
            db.session.execute(text('ALTER TABLE user ADD COLUMN username VARCHAR(100)'))
            db.session.commit()
    if 'audit' in tables:
        audit_cols = {col['name'] for col in inspector.get_columns('audit')}
        if 'reason' not in audit_cols:
            db.session.execute(text("ALTER TABLE audit ADD COLUMN reason VARCHAR(255) DEFAULT ''"))
            db.session.commit()
        if 'details' not in audit_cols:
            db.session.execute(text("ALTER TABLE audit ADD COLUMN details TEXT DEFAULT ''"))
            db.session.commit()
    if 'person' not in tables:
        return
    cols = {col['name'] for col in inspector.get_columns('person')}
    additions = [
        ('address_line1', 'VARCHAR(120) DEFAULT \'\''),
        ('address_line2', 'VARCHAR(120) DEFAULT \'\''),
        ('city', 'VARCHAR(100) DEFAULT \'\''),
        ('state', 'VARCHAR(2) DEFAULT \'\''),
        ('zip_code', 'VARCHAR(20) DEFAULT \'\''),
    ]
    for name, definition in additions:
        if name not in cols:
            db.session.execute(text(f'ALTER TABLE person ADD COLUMN {name} {definition}'))
            db.session.commit()
            cols.add(name)
    if 'address' in cols:
        db.session.execute(text(
            "UPDATE person SET address_line1 = address "
            "WHERE (address_line1 IS NULL OR address_line1 = '') "
            "AND address IS NOT NULL AND address != ''"
        ))
        db.session.commit()
    if 'fund' in tables:
        for old_name, new_name in FUND_ALIASES.items():
            existing_old = db.session.scalar(db.select(Fund).filter_by(name=old_name))
            existing_new = db.session.scalar(db.select(Fund).filter_by(name=new_name))
            if existing_old and not existing_new:
                existing_old.name = new_name
                db.session.commit()
            elif existing_old and existing_new:
                db.session.execute(text('UPDATE gift SET fund_id = :new_id WHERE fund_id = :old_id'), {'new_id': existing_new.id, 'old_id': existing_old.id})
                db.session.delete(existing_old)
                db.session.commit()
        for name in DEFAULT_FUNDS:
            if not db.session.scalar(db.select(Fund.id).filter_by(name=name)):
                db.session.add(Fund(name=name))
        db.session.commit()

def funds_in_display_order(funds):
    rank = {name: index for index, name in enumerate(DEFAULT_FUNDS)}
    return sorted(funds, key=lambda fund: (rank.get(fund.name, len(DEFAULT_FUNDS)), fund.name.lower()))

FUND_CHART_COLORS = {
    'General Fund': '#1b624c',
    'Missions Fund': '#3d9b74',
    'Building Fund': '#c4a35a',
    'Benevolence Fund': '#3d6b8a',
}
FUND_CHART_EXTRA_COLORS = ('#7a5c45', '#5c7a6e', '#8b6b4a', '#2a5a6b', '#6b4f7a', '#4a6b58')

def fund_chart_color(name, index=0):
    return FUND_CHART_COLORS.get(name) or FUND_CHART_EXTRA_COLORS[index % len(FUND_CHART_EXTRA_COLORS)]

def iter_month_starts(start, end):
    cursor = date(start.year, start.month, 1)
    last = date(end.year, end.month, 1)
    while cursor <= last:
        yield cursor
        if cursor.month == 12:
            cursor = date(cursor.year + 1, 1, 1)
        else:
            cursor = date(cursor.year, cursor.month + 1, 1)

def fund_giving_slices(fiscal_year_id=None):
    """Active gift totals by fund for the overview pie chart."""
    amount_col = db.func.coalesce(db.func.sum(Gift.amount), 0)
    query = (
        db.select(Fund.name, amount_col)
        .join(Gift, Gift.fund_id == Fund.id)
        .where(Gift.status == 'active')
        .group_by(Fund.id, Fund.name)
        .having(amount_col > 0)
        .order_by(amount_col.desc(), Fund.name.asc())
    )
    if fiscal_year_id is not None:
        query = query.where(Gift.fiscal_year_id == fiscal_year_id)
    rows = db.session.execute(query).all()
    total = sum((amount or Decimal('0')) for _, amount in rows)
    if total <= 0:
        return [], Decimal('0')
    slices = []
    cursor = 0.0
    for index, (name, amount) in enumerate(rows):
        amount = amount or Decimal('0')
        percent = float(amount / total * 100)
        color = fund_chart_color(name, index)
        start = cursor
        end = cursor + percent
        slices.append({
            'name': name,
            'amount': amount,
            'percent': percent,
            'color': color,
            'start': start,
            'end': end,
        })
        cursor = end
    if slices:
        # Close floating-point gaps so the pie is a full circle.
        slices[-1]['end'] = 100.0
        slices[-1]['percent'] = max(0.0, 100.0 - slices[-1]['start'])
    return slices, total

def monthly_giving_bars(fiscal_year, fund_slices):
    """One stacked bar per month, colored by fund, for the dashboard chart."""
    if fiscal_year is None or not fund_slices:
        return []
    colors = {item['name']: item['color'] for item in fund_slices}
    rank = {name: index for index, name in enumerate(DEFAULT_FUNDS)}
    names = sorted(colors, key=lambda name: (rank.get(name, len(DEFAULT_FUNDS)), name.lower()))
    months = list(iter_month_starts(fiscal_year.start_date, fiscal_year.end_date))
    spans_years = months[0].year != months[-1].year if months else False
    year_col = db.extract('year', Gift.date)
    month_col = db.extract('month', Gift.date)
    rows = db.session.execute(
        db.select(year_col, month_col, Fund.name, db.func.coalesce(db.func.sum(Gift.amount), 0))
        .join(Fund, Gift.fund_id == Fund.id)
        .where(Gift.status == 'active', Gift.fiscal_year_id == fiscal_year.id)
        .group_by(year_col, month_col, Fund.id, Fund.name)
    ).all()
    totals = {}
    for year, month, name, amount in rows:
        totals[(int(year), int(month), name)] = amount or Decimal('0')
    bars = []
    max_total = Decimal('0')
    for month_date in months:
        segments = []
        month_total = Decimal('0')
        for name in names:
            amount = totals.get((month_date.year, month_date.month, name), Decimal('0'))
            if amount <= 0:
                continue
            month_total += amount
            segments.append({'name': name, 'amount': amount, 'color': colors[name]})
        max_total = max(max_total, month_total)
        bars.append({
            'label': month_date.strftime("%b '%y" if spans_years else '%b'),
            'title': month_date.strftime('%B %Y'),
            'total': month_total,
            'segments': segments,
        })
    for bar in bars:
        bar['height'] = float(bar['total'] / max_total * 100) if max_total else 0.0
    return bars

def resolve_fund_name(name):
    cleaned = (name or '').strip()
    if not cleaned:
        raise ValueError('fund is required')
    mapped = FUND_ALIASES.get(cleaned, cleaned)
    fund = db.session.scalar(db.select(Fund).where(Fund.active.is_(True), Fund.name == mapped))
    if fund:
        return fund
    if not mapped.endswith(' Fund'):
        fund = db.session.scalar(db.select(Fund).where(Fund.active.is_(True), Fund.name == f'{mapped} Fund'))
        if fund:
            return fund
    raise ValueError(f'unknown fund "{cleaned}"')

def resolve_import_person(row):
    first = (row.get('first_name') or '').strip()
    last = (row.get('last_name') or '').strip()
    email = (row.get('email') or '').strip().lower()
    if not first or not last:
        raise ValueError('first_name and last_name are required')
    matches = db.session.scalars(
        db.select(Person).where(
            Person.active.is_(True),
            Person.first_name.ilike(first),
            Person.last_name.ilike(last),
        )
    ).all()
    if email:
        email_matches = [person for person in matches if (person.email or '').lower() == email]
        if email_matches:
            matches = email_matches
        elif not matches:
            matches = db.session.scalars(
                db.select(Person).where(Person.active.is_(True), Person.email == email)
            ).all()
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ValueError(f'no active person matched {first} {last}')
    raise ValueError(f'multiple people matched {first} {last}; include email to disambiguate')

def parse_gift_import_row(row, cache):
    person = resolve_import_person(row)
    fund = resolve_fund_name(row.get('fund') or '')
    method = (row.get('method') or '').strip()
    if method not in ('Check', 'Cash', 'Online', 'Other'):
        raise ValueError('method must be Check, Cash, Online, or Other')
    try:
        gift_date = date.fromisoformat((row.get('date') or '').strip())
    except ValueError as exc:
        raise ValueError('date must be YYYY-MM-DD') from exc
    raw_amount = (row.get('amount') or '').strip()
    if not raw_amount:
        raise ValueError('amount is required')
    try:
        amount = parse_amount(raw_amount)
    except ValueError as exc:
        raise ValueError('amount must be a positive number (for example 40.00)') from exc
    year_name = (row.get('fiscal_year') or '').strip()
    if year_name:
        year = cache['years'].get(year_name) or db.session.scalar(db.select(FiscalYear).filter_by(name=year_name))
        if year:
            cache['years'][year_name] = year
    else:
        year = next((item for item in cache['open_years'] if item.start_date <= gift_date <= item.end_date), None)
    if not year:
        raise ValueError('fiscal year not found' if year_name else 'date does not fall in an open fiscal year')
    if not year.is_open:
        raise ValueError(f'fiscal year {year.name} is closed')
    if not (year.start_date <= gift_date <= year.end_date):
        raise ValueError(f'date must fall within fiscal year {year.name}')
    return {
        'person': person,
        'fund': fund,
        'year': year,
        'date': gift_date,
        'amount': amount,
        'method': method,
        'reference': (row.get('reference') or '').strip(),
        'note': (row.get('note') or '').strip(),
    }

def csv_row_blank(row):
    return not any(str(value or '').strip() for value in row.values())

def read_boa_transaction_rows(file_storage):
    raw = file_storage.read()
    if not raw or not (raw.strip() if isinstance(raw, (bytes, str)) else str(raw).strip()):
        raise ValueError('The file is empty.')
    if isinstance(raw, bytes):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeError as exc:
            raise ValueError('The file could not be read. Save it again as a CSV from Excel or Google Sheets.') from exc
    else:
        text = raw
    lines = text.splitlines()
    if not lines:
        raise ValueError('The file is empty.')
    start = 0
    for index, line in enumerate(lines):
        if line.lower().startswith('date,description'):
            start = index
            break
    reader = csv.DictReader(io.StringIO('\n'.join(lines[start:])))
    if not reader.fieldnames or 'Description' not in reader.fieldnames:
        raise ValueError('The CSV must include a Description column (Bank of America export).')
    return list(reader)

def parse_zelle_payer_name(description):
    """Return (first_name, last_name) for an incoming Zelle description, else None."""
    if 'zelle' not in description.lower():
        return None
    if re.search(r'\bpayment\s+to\b', description, re.I):
        return None
    match = re.search(r'\bfrom\s+(.+)$', description, re.I | re.S)
    if not match:
        return None
    rest = match.group(1)
    cut = re.search(r'\s+(?:for\b|Conf(?:#|\b)|CONF(?:#|\b))', rest, re.I)
    name_part = rest[:cut.start()] if cut else rest
    words = [re.sub(r'^[^\w\-]+|[^\w\-]+$', '', word) for word in name_part.split()]
    words = [word for word in words if word]
    if len(words) < 2:
        return None
    first_name, last_name = words[0], words[-1]
    if first_name.lower() in ('maria', 'ma') and len(words) >= 3:
        first_name = f'{words[0]} {words[1]}'
    return first_name, last_name

def parse_zelle_note(description):
    match = re.search(r'\bfor\b(.*)";', description, re.I | re.S)
    if not match:
        return ''
    return match.group(1).strip().strip('"').strip()

def parse_boa_amount(raw):
    text = str(raw or '').strip().replace(',', '').replace('$', '')
    if not text:
        raise ValueError('missing amount')
    value = Decimal(text)
    if not value.is_finite():
        raise ValueError('invalid amount')
    return f'{value.quantize(Decimal("0.01"))}'

def parse_boa_date(raw):
    text = str(raw or '').strip()
    for fmt in ('%m/%d/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError('invalid date')

def extract_people_from_boa_csv(file_storage):
    """Return unique (first_name, last_name) pairs from Bank of America Zelle credits."""
    people = []
    seen = set()
    for row in read_boa_transaction_rows(file_storage):
        parsed = parse_zelle_payer_name((row.get('Description') or '').strip())
        if not parsed:
            continue
        key = (parsed[0].lower(), parsed[1].lower())
        if key in seen:
            continue
        seen.add(key)
        people.append(parsed)
    if not people:
        raise ValueError('No Zelle payer names were found in that file.')
    return people

def extract_gifts_from_boa_csv(file_storage):
    """Return contribution rows from Bank of America Zelle credits."""
    gifts = []
    for row in read_boa_transaction_rows(file_storage):
        description = (row.get('Description') or '').strip()
        parsed = parse_zelle_payer_name(description)
        if not parsed:
            continue
        try:
            gift_date = parse_boa_date(row.get('Date'))
            amount = parse_boa_amount(row.get('Amount'))
        except (ValueError, InvalidOperation):
            continue
        first_name, last_name = parsed
        gifts.append({
            'date': gift_date,
            'first_name': first_name,
            'last_name': last_name,
            'email': '',
            'fund': 'General Fund',
            'amount': amount,
            'method': 'Online',
            'fiscal_year': '',
            'reference': '',
            'note': parse_zelle_note(description),
        })
    if not gifts:
        raise ValueError('No Zelle giving records were found in that file.')
    return gifts

def csv_row_preview(row, fieldnames):
    if not row:
        return ''
    names = fieldnames or list(row.keys())
    return ','.join(str(row.get(name, '') or '').replace('\n', ' ') for name in names)

def read_csv_records(file_storage, required_columns):
    raw = file_storage.read()
    if not raw or not raw.strip():
        raise ValueError('The file is empty.')
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeError as exc:
        raise ValueError('The file could not be read. Save it again as a CSV from Excel or Google Sheets.') from exc
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise ValueError('The file is empty.')
    if ',' not in lines[0]:
        raise ValueError('Columns must be separated by commas. Save the file as CSV and try again.')
    try:
        reader = csv.DictReader(io.StringIO(text))
    except csv.Error as exc:
        raise ValueError('The file is not a valid CSV. Check that commas separate each column.') from exc
    if not reader.fieldnames:
        raise ValueError('The file is missing a header row with column names.')
    original_fields = list(reader.fieldnames)
    normalized = [(name or '').strip() for name in original_fields]
    if len(normalized) != len(set(normalized)):
        raise ValueError('The header row has duplicate column names.')
    if not required_columns.issubset(set(normalized)):
        missing = ', '.join(sorted(required_columns - set(normalized)))
        raise ValueError(f'The CSV is missing required columns: {missing}.')
    key_map = dict(zip(original_fields, normalized))
    try:
        records = []
        for raw_row in reader:
            row = {key_map[key]: (raw_row.get(key) or '') for key in original_fields}
            records.append(row)
    except csv.Error as exc:
        raise ValueError('A row could not be read. Check that commas separate each column.') from exc
    data_rows = [row for row in records if not csv_row_blank(row)]
    if not data_rows:
        raise ValueError('The file has no data rows to import.')
    if len(data_rows) > 10000:
        raise ValueError('CSV exceeds 10,000 rows.')
    return normalized, data_rows

def find_duplicate_person(first, last, email):
    email_l = email.lower()
    matches = db.session.scalars(
        db.select(Person).where(Person.first_name.ilike(first), Person.last_name.ilike(last))
    ).all()
    for person in matches:
        person_email = (person.email or '').strip().lower()
        if email_l and person_email == email_l:
            return person
        if not email_l and not person_email:
            return person
    if email_l:
        return db.session.scalar(db.select(Person).where(Person.email == email))
    return None

PERSON_IMPORT_FIELDS = (
    'first_name', 'last_name', 'member_type', 'email', 'phone',
    'address_line1', 'address_line2', 'city', 'state', 'zip_code',
)
PERSON_IMPORT_FIELD_LABELS = {
    'first_name': 'First name',
    'last_name': 'Last name',
    'member_type': 'Member type',
    'email': 'Email',
    'phone': 'Phone',
    'address_line1': 'Address line 1',
    'address_line2': 'Address line 2',
    'city': 'City',
    'state': 'State',
    'zip_code': 'ZIP code',
}

def pagination_page_items(page, total_pages, window=2):
    """Return page numbers with None placeholders for ellipsis gaps."""
    if total_pages <= 1:
        return [1] if total_pages == 1 else []
    pages = {1, total_pages}
    for number in range(page - window, page + window + 1):
        if 1 <= number <= total_pages:
            pages.add(number)
    ordered = sorted(pages)
    items = []
    previous = None
    for number in ordered:
        if previous is not None and number - previous > 1:
            items.append(None)
        items.append(number)
        previous = number
    return items

PER_PAGE_OPTIONS = (10, 25, 50, 100, 250)

def resolve_per_page(default=25):
    raw = request.args.get('per_page', type=int)
    if raw in PER_PAGE_OPTIONS:
        return raw
    return default if default in PER_PAGE_OPTIONS else 25

def build_pager(total, per_page=None, page_arg='page'):
    total = int(total or 0)
    per_page = resolve_per_page(25 if per_page is None else per_page)
    total_pages = max(1, (total + per_page - 1) // per_page) if total else 1
    page = request.args.get(page_arg, 1, type=int) or 1
    page = min(max(page, 1), total_pages)
    if total == 0:
        range_start = range_end = 0
    else:
        range_start = (page - 1) * per_page + 1
        range_end = min(page * per_page, total)
    return {
        'page': page,
        'per_page': per_page,
        'per_page_options': PER_PAGE_OPTIONS,
        'total': total,
        'total_pages': total_pages,
        'range_start': range_start,
        'range_end': range_end,
        'page_items': pagination_page_items(page, total_pages),
        'offset': (page - 1) * per_page,
    }

def field_is_blank(value):
    return not str(value or '').strip()

def merge_field_value(current, incoming):
    """Portal blank←CSV value; portal value←CSV blank keeps portal; both set←CSV overwrites."""
    if field_is_blank(current) and not field_is_blank(incoming):
        return str(incoming).strip()
    if not field_is_blank(current) and field_is_blank(incoming):
        return current if current is not None else ''
    if not field_is_blank(current) and not field_is_blank(incoming):
        return str(incoming).strip()
    return (current if current is not None else '') if field_is_blank(incoming) else str(incoming).strip()

def find_person_for_people_import(first, last, email):
    """Match an existing person for CSV upsert; raise if the name is ambiguous."""
    email_l = (email or '').strip().lower()
    matches = db.session.scalars(
        db.select(Person).where(Person.first_name.ilike(first), Person.last_name.ilike(last))
    ).all()
    if email_l:
        email_matches = [person for person in matches if (person.email or '').strip().lower() == email_l]
        if len(email_matches) == 1:
            return email_matches[0]
        if len(email_matches) > 1:
            raise ValueError(f'multiple people matched {first} {last} with that email')
        if len(matches) == 1:
            return matches[0]
        if len(matches) > 1:
            raise ValueError(f'multiple people matched {first} {last}; include a unique email to disambiguate')
        return db.session.scalar(db.select(Person).where(db.func.lower(Person.email) == email_l))
    blank_email = [person for person in matches if field_is_blank(person.email)]
    if len(blank_email) == 1:
        return blank_email[0]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError(f'multiple people matched {first} {last}; include email to disambiguate')
    return None

def parse_people_import_fields(row):
    first = (row.get('first_name') or '').strip()
    last = (row.get('last_name') or '').strip()
    kind = (row.get('member_type') or '').strip()
    if not first or not last:
        raise ValueError('first name and last name are required')
    line1, line2, city, state, zip_code = read_person_address(row)
    fields = {
        'first_name': first,
        'last_name': last,
        'member_type': kind,
        'email': (row.get('email') or '').strip(),
        'phone': (row.get('phone') or '').strip(),
        'address_line1': line1,
        'address_line2': line2,
        'city': city,
        'state': state,
        'zip_code': zip_code,
    }
    return fields

def plan_person_merge(person, fields):
    """Return (merged_values, change_dicts) using people-import merge rules."""
    merged = {}
    changes = []
    for name in PERSON_IMPORT_FIELDS:
        current = getattr(person, name) or ''
        incoming = fields.get(name, '')
        if name == 'member_type' and field_is_blank(incoming):
            new_value = current
        else:
            new_value = merge_field_value(current, incoming)
        merged[name] = new_value
        if str(current).strip() != str(new_value).strip():
            changes.append({
                'field': name,
                'label': PERSON_IMPORT_FIELD_LABELS.get(name, name),
                'from': str(current),
                'to': str(new_value),
            })
    if not field_is_blank(fields.get('member_type')) and fields['member_type'] not in ('Member', 'Non-member'):
        raise ValueError('member_type must be Member or Non-member')
    if field_is_blank(merged.get('member_type')):
        raise ValueError('member_type must be Member or Non-member')
    if merged['member_type'] not in ('Member', 'Non-member'):
        raise ValueError('member_type must be Member or Non-member')
    return merged, changes

def apply_person_fields(person, fields):
    for name in PERSON_IMPORT_FIELDS:
        setattr(person, name, fields.get(name) or '')

def people_import_pending_path():
    folder = Path(current_app.instance_path)
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f'people_import_pending_{g.user.id}.json'

def store_people_import_pending(payload):
    people_import_pending_path().write_text(json.dumps(payload), encoding='utf-8')

def load_people_import_pending():
    path = people_import_pending_path()
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        return None

def clear_people_import_pending():
    path = people_import_pending_path()
    if path.exists():
        path.unlink()

def prepare_people_import(fieldnames, records):
    creates, updates, unchanged, failures = [], [], [], []
    seen = set()
    for index, row in enumerate(records, start=2):
        preview = csv_row_preview(row, fieldnames)
        try:
            fields = parse_people_import_fields(row)
            if fields['member_type'] and fields['member_type'] not in ('Member', 'Non-member'):
                raise ValueError('member_type must be Member or Non-member')
            dup_key = (fields['first_name'].lower(), fields['last_name'].lower(), fields['email'].lower())
            if dup_key in seen:
                raise ValueError('duplicate of another row in this file')
            seen.add(dup_key)
            existing = find_person_for_people_import(fields['first_name'], fields['last_name'], fields['email'])
            if existing:
                merged, changes = plan_person_merge(existing, fields)
                item = {
                    'row': index,
                    'person_id': existing.id,
                    'name': existing.name,
                    'fields': merged,
                    'changes': changes,
                    'preview': preview,
                }
                if changes:
                    updates.append(item)
                else:
                    unchanged.append(item)
            else:
                if fields['member_type'] not in ('Member', 'Non-member'):
                    raise ValueError('member_type must be Member or Non-member')
                creates.append({'row': index, 'fields': fields, 'preview': preview})
        except ValueError as exc:
            failures.append({'row': index, 'reason': str(exc), 'preview': preview, 'row_data': row})
    return {
        'creates': creates,
        'updates': updates,
        'unchanged': unchanged,
        'failures': failures,
        'fieldnames': fieldnames,
    }

def commit_people_import(plan):
    successes = 0
    failures = list(plan.get('failures') or [])
    for item in plan.get('creates') or []:
        person = Person()
        apply_person_fields(person, item['fields'])
        db.session.add(person)
        successes += 1
    for item in plan.get('updates') or []:
        person = db.session.get(Person, item['person_id'])
        if person is None:
            failures.append({
                'row': item['row'],
                'reason': 'matched person was deleted before import completed',
                'preview': item.get('preview', ''),
                'row_data': item.get('fields') or {},
            })
            continue
        apply_person_fields(person, item['fields'])
        successes += 1
    successes += len(plan.get('unchanged') or [])
    if successes:
        created = len(plan.get('creates') or [])
        updated = len(plan.get('updates') or [])
        audit(
            'import',
            'person',
            summary=f'{successes} loaded ({created} new, {updated} updated), {len(failures)} failed',
        )
        db.session.commit()
    else:
        db.session.rollback()
    return successes, failures

def find_duplicate_gift(person_id, gift_date, amount, fund_id, method, reference):
    return db.session.scalar(
        db.select(Gift).where(
            Gift.status == 'active',
            Gift.person_id == person_id,
            Gift.date == gift_date,
            Gift.amount == amount,
            Gift.fund_id == fund_id,
            Gift.method == method,
            Gift.reference == reference,
        )
    )

def import_report_path():
    folder = Path(current_app.instance_path)
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f'import_report_{g.user.id}.json'

def store_import_report(report):
    import_report_path().write_text(json.dumps(report), encoding='utf-8')

def load_import_report():
    path = import_report_path()
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError):
        return None

def build_failure_csv(fieldnames, failures):
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=list(fieldnames) + ['failure_reason'], extrasaction='ignore')
    writer.writeheader()
    for item in failures:
        row = dict(item.get('row_data') or {})
        row['failure_reason'] = item['reason']
        writer.writerow({name: row.get(name, '') for name in writer.fieldnames})
    return out.getvalue()

def statement_year_label(gifts):
    year_id = request.args.get('year_id', type=int)
    if year_id:
        fiscal = db.session.get(FiscalYear, year_id)
        if fiscal:
            return fiscal.name
    years = sorted({gift.date.year for gift in gifts})
    if len(years) == 1:
        return str(years[0])
    if len(years) > 1:
        return f'{years[0]}–{years[-1]}'
    return str(date.today().year)

def giving_letter_template_path():
    return Path(current_app.root_path) / 'letter_templates' / 'giving_letter.docx'

def giving_letter_header_path():
    return Path(current_app.root_path) / 'letter_templates' / 'giving_letter_header.png'

def person_statement_address(person):
    lines = person_address_lines(person)
    street = ', '.join(lines[:-1]) if len(lines) > 1 else (lines[0] if lines else '')
    city_state_zip = lines[-1] if lines else ''
    if len(lines) == 1:
        street, city_state_zip = lines[0], ''
    return street, city_state_zip

CHURCH_WEBSITE_URL = 'https://www.greenbrookchurch.org'
CHURCH_WEBSITE_LABEL = 'www.greenbrookchurch.org'
GIVING_LETTER_VERSE = (
    'The point is this: whoever sows sparingly will also reap sparingly, and whoever sows '
    'bountifully will also reap bountifully. Each one must give as he has decided in his heart, '
    'not reluctantly or under compulsion, for God loves a cheerful giver. And God is able to '
    'make all grace abound to you, so that having all sufficiency in all things at all times, '
    'you may abound in every good work.'
)
GIVING_LETTER_VERSE_REF = '2 Corinthians 9:6-8'

def replace_docx_text(docx_path, replacements):
    from docx import Document
    document = Document(str(docx_path))
    for paragraph in document.paragraphs:
        original = paragraph.text
        updated = original
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated == original or not paragraph.runs:
            continue
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ''
    document.save(str(docx_path))

def statement_letter_date(year_label):
    """Giving letters for year Y are dated January 31 of Y+1."""
    digits = re.search(r'(\d{4})', str(year_label or ''))
    year = int(digits.group(1)) if digits else date.today().year
    return f'January 31, {year + 1}'

def _copy_run_font(source_run, target_run):
    """Keep rebuilt body runs at the same face/size as neighboring template text."""
    if source_run is None or target_run is None:
        return
    if source_run.font.name:
        target_run.font.name = source_run.font.name
    if source_run.font.size:
        target_run.font.size = source_run.font.size

def _bold_amount_in_paragraph(paragraph, amount_text):
    display = amount_text if str(amount_text).startswith('$') else f'${amount_text}'
    full = paragraph.text
    if display in full:
        match = display
    elif display.replace('$', '') in full:
        match = display.replace('$', '')
        display = f'${match}' if not match.startswith('$') else match
    else:
        return
    before, after = full.split(match, 1)
    if match != display and before.endswith('$'):
        before = before[:-1]
    sample = next((run for run in paragraph.runs if run.font.size or run.font.name), None)
    for run in paragraph.runs:
        run.text = ''
    if not paragraph.runs:
        paragraph.add_run('')
    paragraph.runs[0].text = before
    paragraph.runs[0].bold = False
    _copy_run_font(sample, paragraph.runs[0])
    amount_run = paragraph.add_run(display)
    amount_run.bold = True
    _copy_run_font(sample, amount_run)
    after_run = paragraph.add_run(after)
    after_run.bold = False
    _copy_run_font(sample, after_run)

def _clear_docx_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = ''
    if not paragraph.runs:
        paragraph.add_run('')

def _remove_docx_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)

def enhance_giving_letter_docx(docx_path, person_name, amount_text, letter_date):
    """Insert letter date, tighten page-1 spacing, bold Dear-name and gift amount."""
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document(str(docx_path))

    # Drop the layered Word letterhead — polish stamps a clean header image.
    # Keep body just under that banner so the header–date gap is not wasted.
    for section in document.sections:
        section.top_margin = Inches(1.5)
        section.bottom_margin = Inches(0.85)
        header = section.header
        for table in list(header.tables):
            _remove_docx_element(table._element)
        for paragraph in header.paragraphs:
            _clear_docx_paragraph(paragraph)

    # Leading tab/blank lines push the date down under the stamped header.
    while document.paragraphs and not document.paragraphs[0].text.strip():
        _remove_docx_element(document.paragraphs[0]._element)

    inserted_date = False
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if not inserted_date and text == person_name:
            date_paragraph = paragraph.insert_paragraph_before(letter_date)
            for run in date_paragraph.runs:
                run.bold = False
            # Keep a blank line between the letter date and the recipient name.
            date_paragraph.paragraph_format.space_after = Pt(14)
            spacer = paragraph.insert_paragraph_before(' ')
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(0)
            inserted_date = True
        if text.startswith('Dear ') and person_name in text:
            for run in paragraph.runs:
                if run.text.strip() == person_name or run.text == person_name:
                    run.bold = True
        if 'amount of' in text.lower() and (amount_text in text or amount_text.replace('$', '') in text):
            _bold_amount_in_paragraph(paragraph, amount_text)

    # Template has two empty paragraphs between Sincerely and Romella; keep one.
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text.strip().startswith('Sincerely'):
            continue
        blanks = []
        for following in paragraphs[index + 1:]:
            if following.text.strip():
                break
            blanks.append(following)
        for blank in blanks[1:]:
            _remove_docx_element(blank._element)
        break

    # Verse is restored during PDF polish so LibreOffice page-2 spill is not required.
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith('The point is this') or '2 Corinthians 9:6-8' in text:
            _remove_docx_element(paragraph._element)
    while document.paragraphs and not document.paragraphs[-1].text.strip():
        _remove_docx_element(document.paragraphs[-1]._element)

    document.save(str(docx_path))

def render_giving_letter_docx(person, year_label, amount_text):
    template = giving_letter_template_path()
    if not template.exists():
        raise FileNotFoundError(f'Missing giving letter template: {template}')
    street, city_state_zip = person_statement_address(person)
    filled = Path(current_app.instance_path) / f'statement_{person.id}_{secrets.token_hex(4)}.docx'
    filled.parent.mkdir(parents=True, exist_ok=True)
    document = MailMerge(str(template))
    document.merge(
        Name=person.name,
        Address=street,
        City_State_Zip=city_state_zip,
        Total=amount_text.replace('$', ''),
    )
    document.write(str(filled))
    # Template hard-codes calendar year as split runs ("202" + "5"); normalize to statement year.
    replace_docx_text(filled, {'2025': str(year_label)})
    enhance_giving_letter_docx(
        filled,
        person_name=person.name,
        amount_text=amount_text if amount_text.startswith('$') else f'${amount_text}',
        letter_date=statement_letter_date(year_label),
    )
    return filled

MACOS_SOFFICE = '/Applications/LibreOffice.app/Contents/MacOS/soffice'

def find_soffice():
    configured = os.environ.get('SOFFICE_PATH')
    if configured:
        return configured
    found = __import__('shutil').which('soffice') or __import__('shutil').which('libreoffice')
    if found:
        return found
    if os.path.exists(MACOS_SOFFICE):
        return MACOS_SOFFICE
    raise RuntimeError(
        'LibreOffice is not installed or could not be found. Install it from libreoffice.org, '
        'or set SOFFICE_PATH to the soffice executable.'
    )

def convert_docx_to_pdf(docx_path):
    out_dir = Path(current_app.instance_path) / 'statement_pdf'
    out_dir.mkdir(parents=True, exist_ok=True)
    result = __import__('subprocess').run(
        [find_soffice(), '--headless', '--nologo', '--nofirststartwizard', '--convert-to', 'pdf', '--outdir', str(out_dir), str(docx_path)],
        capture_output=True,
        text=True,
        check=False,
    )
    pdf_path = out_dir / f'{docx_path.stem}.pdf'
    if result.returncode != 0 or not pdf_path.exists():
        detail = (result.stderr or result.stdout or 'LibreOffice conversion failed').strip()
        raise RuntimeError(detail)
    return polish_giving_letter_pdf(pdf_path)

def draw_letter_header(pdf, page_w, page_h):
    """Draw letterhead artwork and clickable website matching the visible URL."""
    header = giving_letter_header_path()
    draw_h = 0
    if header.exists():
        img = ImageReader(str(header))
        iw, ih = img.getSize()
        draw_h = page_w * (ih / float(iw))
        pdf.setFillColor(colors.white)
        pdf.rect(0, page_h - draw_h - 2, page_w, draw_h + 4, fill=1, stroke=0)
        pdf.drawImage(img, 0, page_h - draw_h, width=page_w, height=draw_h, mask='auto')
        # Clickable area over the visible www.greenbrookchurch.org text (lower-right of banner).
        link_left = page_w * 0.62
        link_right = page_w - 18
        link_bottom = page_h - draw_h + draw_h * 0.08
        link_top = page_h - draw_h + draw_h * 0.42
        pdf.linkURL(
            CHURCH_WEBSITE_URL,
            (link_left, link_bottom, link_right, link_top),
            relative=0,
            thickness=0,
        )
    return draw_h

def draw_giving_letter_verse(pdf, page_w, page_h):
    """Draw the full 2 Corinthians footer verse on page 1."""
    left, right = 54, page_w - 54
    width = right - left
    font, size, leading = 'Times-Italic', 9, 11
    lines = simpleSplit(GIVING_LETTER_VERSE, font, size, width) or [GIVING_LETTER_VERSE]
    verse_top = 28 + len(lines) * leading
    # Cover LibreOffice header/footer fragments that land near the bottom.
    pdf.setFillColor(colors.white)
    pdf.rect(0, 0, page_w, verse_top + 16, fill=1, stroke=0)
    y = verse_top
    pdf.setFillColor(colors.black)
    pdf.setFont(font, size)
    for line in lines:
        pdf.drawString(left, y, line)
        y -= leading
    pdf.setFont('Times-Roman', 9)
    pdf.drawRightString(right, 18, GIVING_LETTER_VERSE_REF)

def polish_giving_letter_pdf(pdf_path):
    """Restore letterhead and scripture LibreOffice drops from the layered Word header."""
    pdf_path = Path(pdf_path)
    reader = PdfReader(str(pdf_path))
    if not reader.pages:
        raise RuntimeError('Letter PDF has no pages')
    page = reader.pages[0]
    page_w = float(page.mediabox.width)
    page_h = float(page.mediabox.height)
    first_page_text = page.extract_text() or ''

    overlay = io.BytesIO()
    c = canvas.Canvas(overlay, pagesize=(page_w, page_h))
    draw_letter_header(c, page_w, page_h)
    # Keep page 1 only; redraw the full verse when LibreOffice spilled it.
    if 'The point is this' not in first_page_text:
        draw_giving_letter_verse(c, page_w, page_h)
    elif GIVING_LETTER_VERSE_REF not in first_page_text:
        c.setFillColor(colors.black)
        c.setFont('Times-Roman', 9)
        c.drawRightString(page_w - 54, 18, GIVING_LETTER_VERSE_REF)
    c.save()
    overlay.seek(0)

    writer = PdfWriter()
    writer.add_page(page).merge_page(PdfReader(overlay).pages[0])
    polished = pdf_path.with_name(pdf_path.stem + '_polished.pdf')
    with polished.open('wb') as handle:
        writer.write(handle)
    return polished

def build_itemized_statement_pdf(person, gifts, church_name, year_label, amount_text):
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=letter, pageCompression=0 if current_app.config.get('TESTING') else 1)
    page_w, page_h = letter
    left, width = 54, 504

    def start_page():
        header_h = draw_letter_header(pdf, page_w, page_h)
        y = page_h - header_h - 28
        pdf.setFillColor(colors.HexColor('#174e3b'))
        pdf.setFont('Helvetica-Bold', 14)
        pdf.drawString(left, y, church_name)
        pdf.setFillColor(colors.black)
        y -= 20
        pdf.setFont('Helvetica', 10)
        pdf.drawString(left, y, f'Itemized contributions for {person.name} — {year_label}')
        y -= 28
        pdf.setFont('Helvetica-Bold', 10)
        pdf.drawString(left, y, 'Date')
        pdf.drawString(140, y, 'Fund')
        pdf.drawRightString(left + width, y, 'Amount')
        return y - 20

    y = start_page()
    pdf.setFont('Helvetica', 10)
    for gift in gifts:
        if y < 72:
            pdf.showPage()
            y = start_page()
            pdf.setFont('Helvetica', 10)
        pdf.drawString(left, y, gift.date.isoformat())
        pdf.drawString(140, y, gift.fund.name[:50])
        pdf.drawRightString(left + width, y, f'${gift.amount:,.2f}')
        y -= 18
    pdf.setFont('Helvetica-Bold', 11)
    pdf.drawRightString(left + width, max(56, y - 10), f'Total: {amount_text}')
    pdf.save(); out.seek(0)
    return out

def merge_pdf_bytes(*pdf_buffers):
    writer = PdfWriter()
    for buffer in pdf_buffers:
        reader = PdfReader(buffer)
        for page in reader.pages:
            writer.add_page(page)
    merged = io.BytesIO()
    writer.write(merged)
    merged.seek(0)
    return merged

def draw_paragraphs(pdf, paragraphs, x, y, width, font='Helvetica', size=10, leading=14, gap=10):
    pdf.setFont(font, size)
    for paragraph in paragraphs:
        lines = simpleSplit(paragraph, font, size, width) or ['']
        for line in lines:
            if y < 72:
                pdf.showPage()
                pdf.setFont(font, size)
                y = 740
            pdf.drawString(x, y, line)
            y -= leading
        y -= gap
    return y

def create_app(test_config=None):
    if os.environ.get('PORTAL_PRODUCTION') == '1' and not os.environ.get('SECRET_KEY') and not test_config:
        raise RuntimeError('SECRET_KEY must be set in production.')
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.environ.get('SECRET_KEY', secrets.token_hex(32)),
        SQLALCHEMY_DATABASE_URI=os.environ.get('DATABASE_URL', 'sqlite:///gbbc_portal.db'),
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        # PythonAnywhere MySQL drops connections idle for 300 seconds.
        SQLALCHEMY_ENGINE_OPTIONS={'pool_recycle': 280, 'pool_pre_ping': True},
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE='Lax',
        SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE', '0') == '1' or os.environ.get('PORTAL_PRODUCTION') == '1',
        PERMANENT_SESSION_LIFETIME=INACTIVITY_TIMEOUT,
        SESSION_REFRESH_EACH_REQUEST=True,
        MAX_CONTENT_LENGTH=8 * 1024 * 1024,
        INACTIVITY_TIMEOUT_SECONDS=INACTIVITY_TIMEOUT_SECONDS,
    )
    if test_config: app.config.update(test_config)
    db.init_app(app)
    with app.app_context():
        try:
            ensure_schema()
        except Exception:
            pass

    @app.after_request
    def response_headers(response):
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['X-Frame-Options'] = 'DENY'
        response.headers['Referrer-Policy'] = 'same-origin'
        response.headers['X-Robots-Tag'] = 'noindex, nofollow, noarchive'
        # Templates use inline scripts and handlers, so 'unsafe-inline' is required until they move to static files.
        response.headers['Content-Security-Policy'] = (
            "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; "
            "img-src 'self' data:; object-src 'none'; base-uri 'self'; form-action 'self'; frame-ancestors 'none'"
        )
        if current_app.config['SESSION_COOKIE_SECURE']:
            response.headers['Strict-Transport-Security'] = 'max-age=31536000'
        if request.path not in ('/login', '/signup'): response.headers['Cache-Control'] = 'no-store'
        return response

    @app.before_request
    def load_user_and_csrf():
        g.user = db.session.get(User, session.get('user_id')) if session.get('user_id') else None
        if g.user and not g.user.active:
            session.clear(); g.user = None
        if g.user:
            now = utcnow()
            last_raw = session.get('last_activity')
            if last_raw:
                try:
                    last_activity = datetime.fromisoformat(last_raw)
                except ValueError:
                    last_activity = None
                if last_activity is not None and (now - last_activity) > INACTIVITY_TIMEOUT:
                    audit(
                        'idle_logout',
                        'user',
                        g.user.id,
                        summary='Signed out after 10 minutes of inactivity',
                    )
                    db.session.commit()
                    session.clear()
                    g.user = None
                    flash('You were signed out after 10 minutes of inactivity.', 'error')
                    if request.endpoint not in ('login', 'signup', 'static', 'session_ping'):
                        next_path = request.path if request.method == 'GET' and request.path.startswith('/') and not request.path.startswith('//') else None
                        return redirect(url_for('login', next=next_path) if next_path else url_for('login'))
            if g.user:
                session['last_activity'] = now.isoformat(timespec='seconds')
        if 'csrf' not in session: session['csrf'] = secrets.token_urlsafe(32)
        if request.method in ('POST', 'PUT', 'PATCH', 'DELETE'):
            if not secrets.compare_digest(session['csrf'], request.form.get('csrf_token', '')): abort(400)
        session.permanent = True

    @app.context_processor
    def context():
        def can(permission): return bool(g.user and permission in PERMISSIONS.get(g.user.role, set()))
        return {
            'current_user': g.user,
            'inactivity_timeout_seconds': current_app.config.get('INACTIVITY_TIMEOUT_SECONDS', INACTIVITY_TIMEOUT_SECONDS),
            'can': can,
            'csrf_token': lambda: session['csrf'],
            'roles': ROLES,
            'signup_roles': SIGNUP_ROLES,
            'today': date.today(),
            'us_states': US_STATES,
        }

    def require(permission):
        def deco(fn):
            @wraps(fn)
            def wrapper(*args, **kwargs):
                if not g.user: return redirect(url_for('login', next=request.path))
                if permission not in PERMISSIONS.get(g.user.role, set()): abort(403)
                return fn(*args, **kwargs)
            return wrapper
        return deco

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if g.user:
            return redirect(default_home_path(g.user))
        if request.method == 'POST':
            identifier = request.form.get('email', '').strip().lower()
            recent = db.session.scalar(db.select(db.func.count(LoginAttempt.id)).where(LoginAttempt.email == identifier, LoginAttempt.attempted_at >= utcnow() - timedelta(minutes=15)))
            if recent >= 5:
                flash('Too many sign-in attempts. Try again in 15 minutes.', 'error')
                return render_template('login.html'), 429
            user = db.session.scalar(db.select(User).where(db.or_(User.email == identifier, User.username == identifier), User.active.is_(True)))
            if user and check_password_hash(user.password_hash, request.form.get('password', '')):
                session.clear()
                session['user_id'] = user.id
                session['csrf'] = secrets.token_urlsafe(32)
                session['last_activity'] = utcnow().isoformat(timespec='seconds')
                user.last_login = utcnow(); g.user = user; audit('login', 'user', user.id); db.session.commit()
                target = request.args.get('next') or default_home_path(user)
                if not is_local_path(target):
                    target = default_home_path(user)
                if target == '/' and 'overview' not in PERMISSIONS.get(user.role, set()):
                    target = default_home_path(user)
                return redirect(target)
            pending = db.session.scalar(
                db.select(SignupRequest).where(SignupRequest.email == identifier, SignupRequest.status == 'pending')
            )
            db.session.add(LoginAttempt(email=identifier)); audit('failed_login', 'user', user.id if user else None); db.session.commit()
            if pending:
                flash('Your account request is still waiting for administrator approval.', 'error')
            else:
                flash('Invalid email or password.', 'error')
        return render_template('login.html')

    @app.route('/signup', methods=['GET', 'POST'])
    def signup():
        if g.user:
            return redirect(default_home_path(g.user))
        form = {
            'first_name': '',
            'last_name': '',
            'email': '',
            'role': '',
        }
        if request.method == 'POST':
            form['first_name'] = (request.form.get('first_name') or '').strip()
            form['last_name'] = (request.form.get('last_name') or '').strip()
            form['email'] = (request.form.get('email') or '').strip().lower()
            form['role'] = (request.form.get('role') or '').strip()
            password = request.form.get('password') or ''
            confirm = request.form.get('confirm_password') or ''
            try:
                if not form['first_name'] or not form['last_name']:
                    raise ValueError('First name and last name are required.')
                if len(form['first_name']) > 32 or len(form['last_name']) > 32:
                    raise ValueError('First name and last name are limited to 32 characters.')
                if not valid_email(form['email']):
                    raise ValueError('Enter a valid email address. This email is your user ID.')
                if form['role'] not in SIGNUP_ROLES:
                    raise ValueError('Choose a role: System Administrator, Finance Staff, or People Viewer.')
                validate_signup_password(password)
                if password != confirm:
                    raise ValueError('Password and confirmation do not match.')
                if db.session.scalar(db.select(User).where(db.func.lower(User.email) == form['email'])):
                    raise ValueError('An account with that email already exists. Sign in instead.')
                existing_pending = db.session.scalar(
                    db.select(SignupRequest).where(SignupRequest.email == form['email'], SignupRequest.status == 'pending')
                )
                if existing_pending:
                    raise ValueError('A signup request for that email is already waiting for approval.')
            except ValueError as exc:
                flash(str(exc), 'error')
                return render_template('signup.html', form=form), 400
            signup_req = SignupRequest(
                first_name=form['first_name'][:32],
                last_name=form['last_name'][:32],
                email=form['email'],
                password_hash=hash_password(password),
                requested_role=form['role'],
                status='pending',
            )
            db.session.add(signup_req)
            db.session.flush()
            audit('signup_request', 'signup', signup_req.id, summary=form['role'])
            db.session.commit()
            mail_status = 'queued'
            mail_error = ''
            try:
                mail_status = notify_signup_received(signup_req)
            except Exception as exc:
                current_app.logger.exception('Failed to send signup received acknowledgment')
                mail_status = 'failed'
                mail_error = str(exc)
            return render_template(
                'signup_pending.html',
                email=form['email'],
                mail_status=mail_status,
                mail_configured=mail_is_configured(),
                mail_error=mail_error,
            )
        return render_template('signup.html', form=form)

    @app.post('/logout')
    def logout():
        if g.user: audit('logout', 'user', g.user.id); db.session.commit()
        session.clear(); return redirect(url_for('login'))

    @app.get('/session/ping')
    def session_ping():
        """Keep the session alive while the browser still has user activity."""
        if not g.user:
            return ('', 401)
        return ('', 204)

    @app.get('/robots.txt')
    def robots_txt():
        return ('User-agent: *\nDisallow: /\n', 200, {'Content-Type': 'text/plain; charset=utf-8'})

    @app.route('/')
    @require('overview')
    def dashboard():
        fy = db.session.scalar(db.select(FiscalYear).order_by(FiscalYear.start_date.desc()))
        active_gifts = db.session.scalars(db.select(Gift).filter_by(status='active').order_by(Gift.date.desc()).limit(5)).all() if 'gift_view' in PERMISSIONS[g.user.role] else []
        total_query = db.select(db.func.sum(Gift.amount)).where(Gift.status == 'active')
        if fy: total_query = total_query.where(Gift.fiscal_year_id == fy.id)
        total = db.session.scalar(total_query) or Decimal('0')
        count = db.session.scalar(db.select(db.func.count(Person.id))) or 0
        fund_slices, fund_total = fund_giving_slices(fy.id if fy else None)
        pie_gradient = ', '.join(
            f"{item['color']} {item['start']:.4f}% {item['end']:.4f}%" for item in fund_slices
        ) if fund_slices else ''
        return render_template(
            'dashboard.html',
            gifts=active_gifts,
            total=total,
            people_count=count,
            fy=fy,
            fund_slices=fund_slices,
            fund_total=fund_total,
            pie_gradient=pie_gradient,
            month_bars=monthly_giving_bars(fy, fund_slices),
        )

    @app.route('/people')
    @require('people_view')
    def people():
        sort_keys = {
            'name': (Person.last_name, Person.first_name),
            'type': Person.member_type,
            'contact': Person.email,
            'address': (Person.city, Person.address_line1, Person.last_name),
            'status': Person.active,
        }
        term = request.args.get('q', '').strip()
        sort = request.args.get('sort', 'name')
        if sort not in sort_keys:
            sort = 'name'
        direction = request.args.get('dir', 'asc')
        if direction not in ('asc', 'desc'):
            direction = 'asc'
        filters = []
        if term:
            filters.append(db.or_(Person.first_name.ilike(f'%{term}%'), Person.last_name.ilike(f'%{term}%'), Person.email.ilike(f'%{term}%')))
        count_query = db.select(db.func.count()).select_from(Person)
        if filters:
            count_query = count_query.where(*filters)
        total = db.session.scalar(count_query) or 0
        pager = build_pager(total)
        query = db.select(Person)
        if filters:
            query = query.where(*filters)
        columns = sort_keys[sort]
        if not isinstance(columns, tuple):
            columns = (columns,)
        order = [col.desc() if direction == 'desc' else col.asc() for col in columns]
        order.append(Person.id.desc() if direction == 'desc' else Person.id.asc())
        people_rows = db.session.scalars(query.order_by(*order).offset(pager['offset']).limit(pager['per_page'])).all()
        return render_template(
            'people.html',
            people=people_rows,
            q=term,
            sort=sort,
            direction=direction,
            pager=pager,
        )

    @app.route('/people/new', methods=['GET', 'POST'])
    @require('people_write')
    def person_new(): return person_editor()

    @app.route('/people/<int:person_id>/edit', methods=['GET', 'POST'])
    @require('people_write')
    def person_edit(person_id): return person_editor(db.get_or_404(Person, person_id))

    @app.post('/people/<int:person_id>/delete')
    @require('people_write')
    def person_delete(person_id):
        person = db.get_or_404(Person, person_id)
        gift_count = db.session.scalar(db.select(db.func.count()).select_from(Gift).where(Gift.person_id == person.id)) or 0
        if gift_count:
            flash('Cannot delete a person who has contribution records. Remove or reassign those gifts first.', 'error')
            return redirect(url_for('people'))
        try:
            reason = require_audit_reason()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('people'))
        name = person.name
        db.session.delete(person)
        audit('delete', 'person', person_id, summary=name, reason=reason)
        db.session.commit()
        flash(f'{name} deleted.', 'success')
        return redirect(url_for('people'))

    def person_editor(person=None):
        if request.method == 'POST':
            first = request.form.get('first_name', '').strip(); last = request.form.get('last_name', '').strip()
            kind = request.form.get('member_type', '')
            if not first or not last or kind not in ('Member', 'Non-member'):
                flash('First name, last name, and member type are required.', 'error')
            else:
                try:
                    line1, line2, city, state, zip_code = read_person_address(request.form)
                    reason = require_audit_reason(required=person is not None)
                except ValueError as exc:
                    flash(str(exc), 'error')
                else:
                    creating = person is None
                    before = None if creating else {
                        'First name': person.first_name, 'Last name': person.last_name, 'Member type': person.member_type,
                        'Email': person.email, 'Phone': person.phone, 'Address line 1': person.address_line1,
                        'Address line 2': person.address_line2, 'City': person.city, 'State': person.state, 'ZIP': person.zip_code,
                    }
                    if creating:
                        person = Person(); db.session.add(person)
                    person.first_name = first; person.last_name = last; person.member_type = kind
                    person.email = request.form.get('email', '').strip(); person.phone = request.form.get('phone', '').strip()
                    person.address_line1 = line1; person.address_line2 = line2
                    person.city = city; person.state = state; person.zip_code = zip_code
                    db.session.flush()
                    if creating:
                        audit('create', 'person', person.id, summary=person.name)
                    else:
                        after = {
                            'First name': person.first_name, 'Last name': person.last_name, 'Member type': person.member_type,
                            'Email': person.email, 'Phone': person.phone, 'Address line 1': person.address_line1,
                            'Address line 2': person.address_line2, 'City': person.city, 'State': person.state, 'ZIP': person.zip_code,
                        }
                        changes = [(label, before[label], after[label]) for label in before]
                        audit('update', 'person', person.id, summary=person.name, reason=reason, changes=changes)
                    db.session.commit()
                    flash('Person saved.', 'success'); return redirect(url_for('people'))
        return render_template('person_form.html', person=person)

    def finish_people_import(plan, file_error=''):
        if file_error:
            successes, failures = 0, []
        else:
            successes, failures = commit_people_import(plan)
        store_import_report({
            'kind': 'people',
            'title': 'People import report',
            'success_count': successes,
            'failed_count': len(failures),
            'file_error': file_error,
            'failures': failures,
            'fieldnames': plan.get('fieldnames') or [],
            'back_endpoint': 'people_import',
            'list_endpoint': 'people',
            'list_label': 'View people',
        })
        clear_people_import_pending()
        return redirect(url_for('import_report'))

    @app.route('/people/import', methods=['GET', 'POST'])
    @require('people_write')
    def people_import():
        if request.method == 'POST':
            file = request.files.get('file')
            if not file or not file.filename.lower().endswith('.csv'):
                flash('Choose a CSV file.', 'error')
            else:
                try:
                    fieldnames, records = read_csv_records(file, {'first_name', 'last_name', 'member_type'})
                except ValueError as exc:
                    return finish_people_import({'fieldnames': [], 'creates': [], 'updates': [], 'unchanged': [], 'failures': []}, file_error=str(exc))
                plan = prepare_people_import(fieldnames, records)
                if plan['updates']:
                    store_people_import_pending(plan)
                    return redirect(url_for('people_import_confirm'))
                return finish_people_import(plan)
        return render_template('import.html')

    @app.route('/people/import/confirm', methods=['GET', 'POST'])
    @require('people_write')
    def people_import_confirm():
        plan = load_people_import_pending()
        if not plan or not plan.get('updates'):
            flash('No pending people import updates to confirm.', 'error')
            return redirect(url_for('people_import'))
        if request.method == 'POST':
            action = (request.form.get('action') or '').strip()
            if action == 'cancel':
                clear_people_import_pending()
                flash('People import canceled. No person records were changed.', 'success')
                return redirect(url_for('people_import'))
            return finish_people_import(plan)
        return render_template(
            'people_import_confirm.html',
            updates=plan['updates'],
            create_count=len(plan.get('creates') or []),
            unchanged_count=len(plan.get('unchanged') or []),
            failed_count=len(plan.get('failures') or []),
        )

    @app.route('/people/template.csv')
    @require('people_write')
    def people_template():
        header = 'first_name,last_name,member_type,email,phone,address_line1,address_line2,city,state,zip_code\n'
        content = io.BytesIO(header.encode('utf-8'))
        return send_file(content, as_attachment=True, download_name='GBBC_people_template.csv', mimetype='text/csv')

    @app.route('/contributions')
    @require('gift_view')
    def contributions():
        sort_keys = {
            'date': Gift.date,
            'contributor': (Person.last_name, Person.first_name),
            'fund': Fund.name,
            'method': Gift.method,
            'amount': Gift.amount,
            'reconciled': Gift.reconciled,
            'status': Gift.status,
        }
        sort = request.args.get('sort', 'date')
        if sort not in sort_keys:
            sort = 'date'
        direction = request.args.get('dir', 'desc' if sort == 'date' else 'asc')
        if direction not in ('asc', 'desc'):
            direction = 'desc' if sort == 'date' else 'asc'
        total = db.session.scalar(db.select(db.func.count()).select_from(Gift)) or 0
        pager = build_pager(total)
        query = db.select(Gift)
        if sort == 'contributor':
            query = query.join(Person, Gift.person_id == Person.id)
        elif sort == 'fund':
            query = query.join(Fund, Gift.fund_id == Fund.id)
        columns = sort_keys[sort]
        if not isinstance(columns, tuple):
            columns = (columns,)
        order = [col.desc() if direction == 'desc' else col.asc() for col in columns]
        order.append(Gift.id.desc() if direction == 'desc' else Gift.id.asc())
        gifts = db.session.scalars(query.order_by(*order).offset(pager['offset']).limit(pager['per_page'])).all()
        return render_template(
            'contributions.html',
            gifts=gifts,
            sort=sort,
            direction=direction,
            pager=pager,
        )

    @app.route('/contributions/import', methods=['GET', 'POST'])
    @require('gift_write')
    def gifts_import():
        if request.method == 'POST':
            file = request.files.get('file')
            if not file or not file.filename.lower().endswith('.csv'):
                flash('Choose a CSV file.', 'error')
            else:
                try:
                    fieldnames, records = read_csv_records(file, {'date', 'first_name', 'last_name', 'fund', 'amount', 'method'})
                except ValueError as exc:
                    store_import_report({
                        'kind': 'contributions',
                        'title': 'Contributions import report',
                        'success_count': 0,
                        'failed_count': 0,
                        'file_error': str(exc),
                        'failures': [],
                        'fieldnames': [],
                        'back_endpoint': 'gifts_import',
                        'list_endpoint': 'contributions',
                        'list_label': 'View contributions',
                    })
                    return redirect(url_for('import_report'))
                cache = {
                    'years': {},
                    'open_years': db.session.scalars(db.select(FiscalYear).where(FiscalYear.is_open.is_(True))).all(),
                }
                successes = 0
                failures = []
                seen = set()
                for index, row in enumerate(records, start=2):
                    preview = csv_row_preview(row, fieldnames)
                    try:
                        item = parse_gift_import_row(row, cache)
                        dup_key = (
                            item['person'].id, item['date'].isoformat(), str(item['amount']),
                            item['fund'].id, item['method'], item['reference'],
                        )
                        if dup_key in seen:
                            raise ValueError('duplicate of another row in this file')
                        if find_duplicate_gift(item['person'].id, item['date'], item['amount'], item['fund'].id, item['method'], item['reference']):
                            raise ValueError('duplicate of an existing contribution')
                        seen.add(dup_key)
                        db.session.add(Gift(
                            person=item['person'], fund=item['fund'], fiscal_year=item['year'],
                            date=item['date'], amount=item['amount'], method=item['method'],
                            reference=item['reference'], note=item['note'], entered_by_id=g.user.id,
                        ))
                        successes += 1
                    except ValueError as exc:
                        failures.append({'row': index, 'reason': str(exc), 'preview': preview, 'row_data': row})
                if successes:
                    audit('import', 'gift', summary=f'{successes} loaded, {len(failures)} failed')
                    db.session.commit()
                else:
                    db.session.rollback()
                store_import_report({
                    'kind': 'contributions',
                    'title': 'Contributions import report',
                    'success_count': successes,
                    'failed_count': len(failures),
                    'file_error': '',
                    'failures': failures,
                    'fieldnames': fieldnames,
                    'back_endpoint': 'gifts_import',
                    'list_endpoint': 'contributions',
                    'list_label': 'View contributions',
                })
                return redirect(url_for('import_report'))
        return render_template('contributions_import.html')

    @app.route('/import/report')
    def import_report():
        if not g.user:
            return redirect(url_for('login', next=request.path))
        report = load_import_report()
        if not report:
            flash('No recent import report is available.', 'error')
            return redirect(url_for('dashboard'))
        if report.get('kind') == 'people' and 'people_write' not in PERMISSIONS[g.user.role]:
            abort(403)
        if report.get('kind') == 'contributions' and 'gift_write' not in PERMISSIONS[g.user.role]:
            abort(403)
        return render_template('import_report.html', report=report)

    @app.route('/import/report/failures.csv')
    def import_report_failures():
        if not g.user:
            return redirect(url_for('login', next=request.path))
        report = load_import_report()
        if not report or not report.get('failures'):
            abort(404)
        if report.get('kind') == 'people' and 'people_write' not in PERMISSIONS[g.user.role]:
            abort(403)
        if report.get('kind') == 'contributions' and 'gift_write' not in PERMISSIONS[g.user.role]:
            abort(403)
        content = build_failure_csv(report.get('fieldnames') or [], report['failures'])
        return send_file(
            io.BytesIO(content.encode('utf-8')),
            as_attachment=True,
            download_name=f"GBBC_{report.get('kind', 'import')}_failed_rows.csv",
            mimetype='text/csv',
        )

    @app.route('/contributions/template.csv')
    @require('gift_write')
    def gifts_template():
        header = 'date,first_name,last_name,email,fund,amount,method,fiscal_year,reference,note\n'
        example = '2026-09-01,Ada,Member,ada@example.org,General Fund,25.00,Check,2026,1001,\n'
        content = io.BytesIO((header + example).encode('utf-8'))
        return send_file(content, as_attachment=True, download_name='GBBC_contributions_template.csv', mimetype='text/csv')

    @app.get('/people/search')
    @require('gift_write')
    def people_search():
        """Return active people whose names match the query (min 2 characters).

        Matches any substring across first name, last name, and the combined
        full name so middle names/initials stored in either field are found.
        Multi-word queries require every token to appear somewhere in the name.
        """
        raw = (request.args.get('q') or '').strip()
        if len(raw) < 2:
            return jsonify([])
        tokens = [token for token in re.split(r'\s+', raw) if token]
        if not tokens:
            return jsonify([])
        full_name = db.func.concat(Person.first_name, ' ', Person.last_name)
        query = db.select(Person).where(Person.active.is_(True))
        for token in tokens:
            pattern = f'%{token}%'
            query = query.where(
                db.or_(
                    Person.first_name.ilike(pattern),
                    Person.last_name.ilike(pattern),
                    full_name.ilike(pattern),
                )
            )
        people = db.session.scalars(query.order_by(Person.last_name, Person.first_name).limit(40)).all()
        return jsonify([
            {
                'id': person.id,
                'name': person.name,
                'first_name': person.first_name,
                'last_name': person.last_name,
                'member_type': person.member_type,
            }
            for person in people
        ])

    @app.route('/contributions/new', methods=['GET', 'POST'])
    @require('gift_write')
    def gift_new(): return gift_editor()

    @app.route('/contributions/<int:gift_id>/edit', methods=['GET', 'POST'])
    @require('gift_write')
    def gift_edit(gift_id): return gift_editor(db.get_or_404(Gift, gift_id))

    def gift_editor(gift=None):
        years = db.session.scalars(db.select(FiscalYear).order_by(FiscalYear.start_date.desc())).all()
        funds = funds_in_display_order(db.session.scalars(db.select(Fund).filter_by(active=True)).all())
        default_fund = next((fund for fund in funds if fund.name == 'General Fund'), funds[0] if funds else None)
        selected_person = gift.person if gift else None
        if request.method == 'POST':
            try:
                year = db.session.get(FiscalYear, int(request.form['fiscal_year_id']))
                person = db.session.get(Person, int(request.form['person_id']))
                fund = db.session.get(Fund, int(request.form['fund_id']))
                gift_date = date.fromisoformat(request.form['date']); amount = parse_amount(request.form['amount'])
                if not year or not person or not fund or not year.is_open: raise ValueError('Choose an active person, fund, and open fiscal year.')
                if not person.active and (gift is None or gift.person_id != person.id):
                    raise ValueError('Choose an active person, fund, and open fiscal year.')
                if not (year.start_date <= gift_date <= year.end_date): raise ValueError('Date must fall within the selected fiscal year.')
                if gift and not gift.fiscal_year.is_open: raise ValueError('Closed-year gifts cannot be edited.')
                reason = require_audit_reason(required=gift is not None)
                creating = gift is None
                before = None
                if not creating:
                    before = {
                        'Contributor': gift.person.name, 'Fund': gift.fund.name, 'Amount': gift.amount,
                        'Date': gift.date, 'Method': gift.method, 'Reference': gift.reference, 'Note': gift.note,
                        'Fiscal year': gift.fiscal_year.name,
                    }
                if creating:
                    gift = Gift(entered_by_id=g.user.id); db.session.add(gift)
                gift.person = person; gift.fund = fund; gift.fiscal_year = year; gift.date = gift_date; gift.amount = amount
                gift.method = request.form.get('method', '').strip(); gift.reference = request.form.get('reference', '').strip()
                gift.note = request.form.get('note', '').strip()
                if gift.method not in ('Check', 'Cash', 'Online', 'Other'): raise ValueError('Choose a payment method.')
                db.session.flush()
                if creating:
                    note_text = gift.note or '(empty)'
                    audit(
                        'create',
                        'gift',
                        gift.id,
                        summary=f'{person.name} · {fund.name} · ${amount} · {note_text}',
                        details=f'Contributor: {person.name}\nFund: {fund.name}\nAmount: {amount}\nNote: {note_text}',
                    )
                else:
                    after = {
                        'Contributor': gift.person.name, 'Fund': gift.fund.name, 'Amount': gift.amount,
                        'Date': gift.date, 'Method': gift.method, 'Reference': gift.reference, 'Note': gift.note,
                        'Fiscal year': gift.fiscal_year.name,
                    }
                    changes = [(label, before[label], after[label]) for label in before]
                    audit('update', 'gift', gift.id, summary=person.name, reason=reason, changes=changes)
                db.session.commit()
                flash('Contribution saved.', 'success'); return redirect(url_for('contributions'))
            except (ValueError, KeyError) as exc:
                flash(str(exc), 'error')
                person_id = request.form.get('person_id', type=int)
                selected_person = db.session.get(Person, person_id) if person_id else selected_person
        return render_template(
            'gift_form.html',
            gift=gift,
            years=years,
            funds=funds,
            selected_person=selected_person,
            default_fund_id=default_fund.id if default_fund else None,
            people_search_url=url_for('people_search'),
        )

    @app.post('/contributions/<int:gift_id>/void')
    @require('gift_write')
    def gift_void(gift_id):
        gift = db.get_or_404(Gift, gift_id)
        if not gift.fiscal_year.is_open: abort(403)
        try:
            reason = require_audit_reason()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('contributions'))
        gift.status = 'void'
        audit('void', 'gift', gift.id, summary=gift.person.name, reason=reason, changes=[('Status', 'active', 'void')])
        db.session.commit(); flash('Contribution voided.', 'success')
        return redirect(url_for('contributions'))

    @app.post('/contributions/<int:gift_id>/delete')
    @require('gift_write')
    def gift_delete(gift_id):
        gift = db.get_or_404(Gift, gift_id)
        if not gift.fiscal_year.is_open:
            flash('Closed-year contributions cannot be deleted.', 'error')
            return redirect(url_for('contributions'))
        try:
            reason = require_audit_reason()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('contributions'))
        summary = f'{gift.person.name} ${gift.amount} on {gift.date.isoformat()}'
        details = f'Contributor: {gift.person.name}\nFund: {gift.fund.name}\nAmount: {gift.amount}\nNote: {gift.note or "(empty)"}'
        db.session.delete(gift)
        audit('delete', 'gift', gift_id, summary=summary, reason=reason, details=details)
        db.session.commit()
        flash('Contribution deleted.', 'success')
        return redirect(url_for('contributions'))

    @app.post('/contributions/<int:gift_id>/reconcile')
    @require('gift_write')
    def gift_reconcile(gift_id):
        gift = db.get_or_404(Gift, gift_id)
        if gift.status != 'active' or not gift.fiscal_year.is_open: abort(403)
        try:
            reason = require_audit_reason()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('contributions'))
        old = gift.reconciled
        gift.reconciled = not gift.reconciled
        audit(
            'reconcile', 'gift', gift.id, summary=gift.person.name, reason=reason,
            changes=[('Reconciled', old, gift.reconciled)],
        )
        db.session.commit()
        flash('Reconciliation status updated.', 'success')
        return redirect(url_for('contributions'))

    def filtered_gifts():
        year_id = request.args.get('year_id', type=int)
        person_id = request.args.get('person_id', type=int)
        query = db.select(Gift).where(Gift.status == 'active')
        if year_id: query = query.where(Gift.fiscal_year_id == year_id)
        if person_id: query = query.where(Gift.person_id == person_id)
        for field, bound in (('start_date', 'start'), ('end_date', 'end')):
            raw = request.args.get(field)
            if raw:
                try:
                    parsed = date.fromisoformat(raw)
                except ValueError:
                    abort(400)
                query = query.where(Gift.date >= parsed if bound == 'start' else Gift.date <= parsed)
        return db.session.scalars(query.order_by(Gift.date, Gift.id)).all()

    def report_gifts():
        gifts = filtered_gifts()
        minimum = request.args.get('minimum_total', '').strip()
        if not minimum: return gifts
        try:
            threshold = Decimal(minimum)
            if threshold < 0: raise ValueError
        except (InvalidOperation, ValueError): abort(400)
        person_totals = {}
        for gift in gifts: person_totals[gift.person_id] = person_totals.get(gift.person_id, Decimal('0')) + gift.amount
        return [gift for gift in gifts if person_totals[gift.person_id] >= threshold]

    @app.route('/reports')
    @require('reports')
    def reports():
        if request.args.get('person_id') and 'detail_reports' not in PERMISSIONS[g.user.role]: abort(403)
        if request.args.get('minimum_total') and 'detail_reports' not in PERMISSIONS[g.user.role]: abort(403)
        years = db.session.scalars(db.select(FiscalYear).order_by(FiscalYear.start_date.desc())).all()
        role_perms = PERMISSIONS.get(g.user.role, set())
        can_detail = 'detail_reports' in role_perms
        can_statements = 'statements' in role_perms
        persons = db.session.scalars(db.select(Person).order_by(Person.last_name, Person.first_name)).all() if (can_detail or can_statements) else []
        gifts = report_gifts()
        total = sum((x.amount for x in gifts), Decimal('0'))
        funds = {}
        for gift in gifts: funds[gift.fund.name] = funds.get(gift.fund.name, Decimal('0')) + gift.amount
        monthly = {}
        for gift in gifts:
            period = gift.date.strftime('%b %Y')
            monthly[period] = monthly.get(period, Decimal('0')) + gift.amount
        detail_gifts = gifts if can_detail else []
        pager = build_pager(len(detail_gifts)) if can_detail else None
        page_gifts = detail_gifts[pager['offset']:pager['offset'] + pager['per_page']] if pager else []
        filter_args = {
            'year_id': request.args.get('year_id') or None,
            'start_date': request.args.get('start_date') or None,
            'end_date': request.args.get('end_date') or None,
            'minimum_total': request.args.get('minimum_total') or None,
            'person_id': request.args.get('person_id') or None,
            'per_page': pager['per_page'] if pager else None,
        }
        audit('view', 'report', summary='aggregate' if not persons else 'detail'); db.session.commit()
        return render_template(
            'reports.html',
            years=years,
            people=persons,
            gifts=page_gifts,
            total=total,
            funds=funds,
            monthly=monthly,
            selected_year=request.args.get('year_id', type=int),
            pager=pager,
            filter_args=filter_args,
        )

    @app.route('/reports/export.xlsx')
    @require('detail_reports')
    def report_xlsx():
        gifts = report_gifts(); wb = Workbook(); ws = wb.active; ws.title = 'Giving'
        ws.append(['Date', 'Contributor', 'Address line 1', 'Address line 2', 'City', 'State', 'ZIP', 'Fund', 'Payment method', 'Reference', 'Amount'])
        for gift in gifts:
            person = gift.person
            ws.append([
                gift.date.isoformat(), excel_safe(person.name),
                excel_safe(person.address_line1 or ''), excel_safe(person.address_line2 or ''),
                excel_safe(person.city or ''), excel_safe(person.state or ''), excel_safe(person.zip_code or ''),
                excel_safe(gift.fund.name), excel_safe(gift.method), excel_safe(gift.reference), float(gift.amount),
            ])
        ws.append(['', '', '', '', '', '', '', '', '', 'Total', float(sum((x.amount for x in gifts), Decimal('0')))])
        out = io.BytesIO(); wb.save(out); out.seek(0)
        audit('export', 'report', summary='xlsx'); db.session.commit()
        return send_file(out, download_name='GBBC_giving_report.xlsx', as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.route('/reports/statements/<int:person_id>.pdf')
    @require('statements')
    def statement(person_id):
        person = db.get_or_404(Person, person_id)
        gifts = [x for x in filtered_gifts() if x.person_id == person_id]
        church_setting = db.session.get(Setting, 'church_name')
        church_name = church_setting.value if church_setting else 'Green Brook Baptist Church'
        year_label = statement_year_label(gifts)
        total = sum((x.amount for x in gifts), Decimal('0'))
        amount_text = f'${total:,.2f}'
        docx_path = pdf_path = None
        try:
            docx_path = render_giving_letter_docx(person, year_label, amount_text)
            pdf_path = convert_docx_to_pdf(docx_path)
            letter_pdf = io.BytesIO(pdf_path.read_bytes())
            itemized_pdf = build_itemized_statement_pdf(person, gifts, church_name, year_label, amount_text)
            out = merge_pdf_bytes(letter_pdf, itemized_pdf)
        except Exception as exc:
            current_app.logger.exception('Giving letter template rendering failed')
            flash(f'Could not generate the giving letter from the Word template: {exc}', 'error')
            return redirect(url_for('reports'))
        finally:
            for path in (docx_path, pdf_path):
                if path and Path(path).exists():
                    try:
                        Path(path).unlink()
                    except OSError:
                        pass
        audit('export', 'statement', person_id); db.session.commit()
        return send_file(out, download_name=f'GBBC_statement_{person_id}.pdf', as_attachment=True, mimetype='application/pdf')

    def _handle_signup_decision():
        signup_id = request.form.get('signup_id', type=int)
        action = (request.form.get('action') or '').strip()
        signup_req = db.get_or_404(SignupRequest, signup_id)
        if signup_req.status != 'pending':
            flash('That signup request was already reviewed.', 'error')
            return redirect(url_for('users'))
        if action not in ('approve', 'deny'):
            abort(400)
        try:
            reason = require_audit_reason()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('users'))
        if action == 'approve':
            if signup_req.requested_role not in SIGNUP_ROLES:
                flash('This request has an invalid role and cannot be approved.', 'error')
                return redirect(url_for('users'))
            if db.session.scalar(db.select(User).where(db.func.lower(User.email) == signup_req.email.lower())):
                flash('A user with that email already exists. Deny this request or resolve the conflict.', 'error')
                return redirect(url_for('users'))
            user = User(
                first_name=signup_req.first_name,
                last_name=signup_req.last_name,
                email=signup_req.email.lower(),
                phone='',
                role=signup_req.requested_role,
                password_hash=signup_req.password_hash,
                active=True,
            )
            db.session.add(user)
            db.session.flush()
            signup_req.status = 'approved'
            signup_req.reviewed_at = utcnow()
            signup_req.reviewed_by_id = g.user.id
            audit(
                'approve', 'signup', signup_req.id,
                summary=f'{signup_req.email} as {signup_req.requested_role}',
                reason=reason,
                changes=[('Status', 'pending', 'approved'), ('Role', '', signup_req.requested_role)],
            )
            db.session.commit()
            try:
                notify_signup_decision(signup_req, approved=True, reviewer=g.user)
            except Exception:
                current_app.logger.exception('Failed to notify signup approval')
                flash('Account approved, but the notification email could not be sent.', 'error')
                return redirect(url_for('users'))
            flash(f'Approved {signup_req.name}. A notice was sent from your email address.', 'success')
        else:
            signup_req.status = 'denied'
            signup_req.reviewed_at = utcnow()
            signup_req.reviewed_by_id = g.user.id
            audit(
                'deny', 'signup', signup_req.id, summary=signup_req.email, reason=reason,
                changes=[('Status', 'pending', 'denied')],
            )
            db.session.commit()
            try:
                notify_signup_decision(signup_req, approved=False, reviewer=g.user)
            except Exception:
                current_app.logger.exception('Failed to notify signup denial')
                flash('Request denied, but the notification email could not be sent.', 'error')
                return redirect(url_for('users'))
            flash(f'Denied {signup_req.name}. A notice was sent from your email address.', 'success')
        return redirect(url_for('users'))

    @app.route('/admin/users', methods=['GET', 'POST'])
    @require('users')
    def users():
        if request.method == 'POST':
            if request.form.get('signup_id'):
                return _handle_signup_decision()
            first = request.form.get('first_name', '').strip(); last = request.form.get('last_name', '').strip()
            email = request.form.get('email', '').strip().lower(); phone = request.form.get('phone', '').strip()
            role = request.form.get('role', ''); password = request.form.get('password', '')
            username = request.form.get('username', '').strip().lower() or None
            if not all((first, last, phone, password)) or len(password) < 12 or not valid_email(email) or role not in ROLES:
                flash('First name, last name, valid email, phone, role, and password are required.', 'error')
            elif db.session.scalar(db.select(User).filter_by(email=email)):
                flash('That email address is already used.', 'error')
            elif username and db.session.scalar(db.select(User).filter_by(username=username)):
                flash('That username is already used.', 'error')
            elif role == 'Super Admin' and g.user.role != 'Super Admin':
                abort(403)
            else:
                user = User(first_name=first, last_name=last, email=email, username=username, phone=phone, role=role, password_hash=hash_password(password))
                db.session.add(user); db.session.flush(); audit('create', 'user', user.id, summary=user.email); db.session.commit()
                flash('User created.', 'success'); return redirect(url_for('users'))
        all_users = db.session.scalars(db.select(User).order_by(User.last_name)).all()
        pending = db.session.scalars(
            db.select(SignupRequest).where(SignupRequest.status == 'pending').order_by(SignupRequest.created_at.asc())
        ).all()
        recent = db.session.scalars(
            db.select(SignupRequest).where(SignupRequest.status != 'pending').order_by(SignupRequest.reviewed_at.desc()).limit(25)
        ).all()
        return render_template('users.html', users=all_users, pending=pending, recent=recent)

    @app.post('/admin/users/<int:user_id>/toggle')
    @require('users')
    def user_toggle(user_id):
        user = db.get_or_404(User, user_id)
        if user.id == g.user.id or (user.role == 'Super Admin' and g.user.role != 'Super Admin'): abort(403)
        if user.role == 'Super Admin' and user.active:
            count = db.session.scalar(db.select(db.func.count(User.id)).filter_by(role='Super Admin', active=True))
            if count <= 1: abort(403)
        try:
            reason = require_audit_reason()
        except ValueError as exc:
            flash(str(exc), 'error')
            return redirect(url_for('users'))
        old = user.active
        user.active = not user.active
        audit(
            'status', 'user', user.id, summary=user.email, reason=reason,
            changes=[('Active', old, user.active)],
        )
        db.session.commit()
        return redirect(url_for('users'))

    @app.route('/admin/users/<int:user_id>/edit', methods=['GET', 'POST'])
    @require('users')
    def user_edit(user_id):
        user = db.get_or_404(User, user_id)
        if user.role == 'Super Admin' and g.user.role != 'Super Admin': abort(403)
        if request.method == 'POST':
            first = request.form.get('first_name', '').strip()
            last = request.form.get('last_name', '').strip()
            email = request.form.get('email', '').strip().lower()
            username = request.form.get('username', '').strip().lower() or None
            phone = request.form.get('phone', '').strip()
            role = request.form.get('role', '')
            other = db.session.scalar(db.select(User).where(User.email == email, User.id != user.id))
            other_username = db.session.scalar(db.select(User).where(User.username == username, User.id != user.id)) if username else None
            try:
                reason = require_audit_reason()
            except ValueError as exc:
                flash(str(exc), 'error')
                return render_template('user_edit.html', user=user)
            if not first or not last or not phone or not valid_email(email) or role not in ROLES:
                flash('First name, last name, valid email, phone, and role are required.', 'error')
            elif other:
                flash('That email address is already used.', 'error')
            elif other_username:
                flash('That username is already used.', 'error')
            elif role == 'Super Admin' and g.user.role != 'Super Admin':
                abort(403)
            elif user.role == 'Super Admin' and role != 'Super Admin' and db.session.scalar(db.select(db.func.count(User.id)).filter_by(role='Super Admin', active=True)) <= 1:
                flash('The last active Super Admin cannot change role.', 'error')
            else:
                before = {
                    'First name': user.first_name, 'Last name': user.last_name, 'Email': user.email,
                    'Username': user.username or '', 'Phone': user.phone, 'Role': user.role,
                }
                user.first_name = first; user.last_name = last; user.email = email; user.username = username; user.phone = phone; user.role = role
                password = request.form.get('password', '')
                password_changed = False
                if password:
                    if len(password) < 12:
                        flash('A replacement password must contain at least 12 characters.', 'error')
                        return render_template('user_edit.html', user=user)
                    user.password_hash = hash_password(password)
                    password_changed = True
                after = {
                    'First name': user.first_name, 'Last name': user.last_name, 'Email': user.email,
                    'Username': user.username or '', 'Phone': user.phone, 'Role': user.role,
                }
                changes = [(label, before[label], after[label]) for label in before]
                if password_changed:
                    changes.append(('Password', '(hidden)', '(updated)'))
                audit('update', 'user', user.id, summary=user.email, reason=reason, changes=changes)
                db.session.commit()
                flash('User updated.', 'success'); return redirect(url_for('users'))
        return render_template('user_edit.html', user=user)

    @app.route('/admin/roles')
    @require('roles')
    def role_matrix(): return render_template('roles.html', matrix=PAGE_MATRIX, permissions=PERMISSIONS)

    @app.route('/admin/settings', methods=['GET', 'POST'])
    @require('settings')
    def settings_page():
        if request.method == 'POST':
            action = (request.form.get('action') or 'save').strip()
            if action == 'test_email':
                to_addr = (request.form.get('test_email_to') or g.user.email or '').strip().lower()
                if not valid_email(to_addr):
                    flash('Enter a valid address for the test email.', 'error')
                    return redirect(url_for('settings_page'))
                if not mail_is_configured():
                    flash('Configure an SMTP host before sending a test email.', 'error')
                    return redirect(url_for('settings_page'))
                try:
                    send_portal_email(
                        to_addr,
                        'GBBC Finance mail test',
                        (
                            f'This is a test message from the GBBC Finance portal.\n\n'
                            f'Sent by {g.user.name} at {utcnow().isoformat(sep=" ", timespec="seconds")} UTC.\n'
                        ),
                        outbox_name='mail_test',
                    )
                    if to_addr.endswith('@example.org') or to_addr.endswith('@example.com'):
                        flash(
                            f'Test email was accepted by SMTP for {to_addr}, but @example.org is not a real inbox. '
                            f'Send the next test to your Gmail address.',
                            'success',
                        )
                    else:
                        flash(f'Test email sent to {to_addr}.', 'success')
                except Exception as exc:
                    current_app.logger.exception('Mail test failed')
                    flash(f'Test email failed: {exc}', 'error')
                return redirect(url_for('settings_page'))

            church_name = request.form.get('church_name', '').strip()
            church_address = request.form.get('church_address', '').strip()
            if not church_name or not church_address:
                flash('Church name and mailing address are required.', 'error')
                return redirect(url_for('settings_page'))
            smtp_port = request.form.get('smtp_port', '').strip() or '587'
            if not smtp_port.isdigit():
                flash('SMTP port must be a number (for example 587 or 465).', 'error')
                return redirect(url_for('settings_page'))
            try:
                reason = require_audit_reason()
            except ValueError as exc:
                flash(str(exc), 'error')
                return redirect(url_for('settings_page'))
            before = {item.key: item.value for item in db.session.scalars(db.select(Setting)).all()}
            mail_from = request.form.get('mail_from', '').strip()
            smtp_host = request.form.get('smtp_host', '').strip()
            smtp_username = request.form.get('smtp_username', '').strip()
            encryption = (request.form.get('smtp_encryption') or 'starttls').strip().lower()
            if encryption not in ('starttls', 'ssl', 'none'):
                encryption = 'starttls'
            new_password = request.form.get('smtp_password', '')
            set_setting('church_name', church_name)
            set_setting('church_address', church_address)
            set_setting('mail_from', mail_from)
            set_setting('smtp_host', smtp_host)
            set_setting('smtp_port', smtp_port)
            set_setting('smtp_username', smtp_username)
            set_setting('smtp_encryption', encryption)
            if new_password.strip():
                set_setting('smtp_password', new_password)
            elif request.form.get('clear_smtp_password') == '1':
                set_setting('smtp_password', '')
            after_keys = {
                'church_name': church_name,
                'church_address': church_address,
                'mail_from': mail_from,
                'smtp_host': smtp_host,
                'smtp_port': smtp_port,
                'smtp_username': smtp_username,
                'smtp_encryption': encryption,
            }
            changes = [(key, before.get(key, ''), new_val) for key, new_val in after_keys.items()]
            if new_password.strip() or request.form.get('clear_smtp_password') == '1':
                changes.append(('smtp_password', '(hidden)' if before.get('smtp_password') else '(empty)', '(updated)' if new_password.strip() else '(cleared)'))
            audit('update', 'settings', summary='website settings', reason=reason, changes=changes)
            db.session.commit()
            flash('Website settings saved.', 'success')
            return redirect(url_for('settings_page'))
        values = {item.key: item.value for item in db.session.scalars(db.select(Setting)).all()}
        return render_template(
            'settings.html',
            values=values,
            mail_configured=mail_is_configured(),
            smtp_password_set=bool(values.get('smtp_password')),
        )

    @app.route('/admin/fiscal', methods=['GET', 'POST'])
    @require('fiscal')
    def fiscal():
        if request.method == 'POST':
            action = request.form.get('action')
            if action == 'create':
                try:
                    start = date.fromisoformat(request.form['start_date']); end = date.fromisoformat(request.form['end_date'])
                    if start >= end: raise ValueError('End date must follow start date.')
                    for existing in db.session.scalars(db.select(FiscalYear)).all():
                        if start <= existing.end_date and end >= existing.start_date: raise ValueError('Fiscal years cannot overlap.')
                    year = FiscalYear(name=request.form['name'].strip(), start_date=start, end_date=end)
                    db.session.add(year); db.session.flush(); audit('create', 'fiscal_year', year.id, summary=year.name); db.session.commit()
                except (ValueError, KeyError) as exc: flash(str(exc), 'error')
            elif action == 'toggle':
                try:
                    reason = require_audit_reason()
                except ValueError as exc:
                    flash(str(exc), 'error')
                    return redirect(url_for('fiscal'))
                year = db.get_or_404(FiscalYear, request.form.get('year_id', type=int))
                old = year.is_open
                year.is_open = not year.is_open
                audit(
                    'status', 'fiscal_year', year.id, summary=year.name, reason=reason,
                    changes=[('Open', old, year.is_open)],
                )
                db.session.commit()
            return redirect(url_for('fiscal'))
        years = db.session.scalars(db.select(FiscalYear).order_by(FiscalYear.start_date.desc())).all()
        return render_template('fiscal.html', years=years)

    @app.route('/utilities')
    @require('utilities')
    def utilities():
        return render_template('utilities.html')

    @app.route('/utilities/boa-people', methods=['GET', 'POST'])
    @require('utilities')
    def utilities_boa_people():
        if request.method == 'POST':
            file = request.files.get('file')
            if not file or not file.filename.lower().endswith('.csv'):
                flash('Choose a Bank of America CSV file.', 'error')
            else:
                try:
                    people = extract_people_from_boa_csv(file)
                except ValueError as exc:
                    flash(str(exc), 'error')
                else:
                    output = io.StringIO()
                    writer = csv.writer(output)
                    writer.writerow(['first_name', 'last_name'])
                    writer.writerows(people)
                    content = io.BytesIO(output.getvalue().encode('utf-8'))
                    audit('export', 'utility', summary=f'BoA Zelle people list ({len(people)})')
                    db.session.commit()
                    return send_file(
                        content,
                        as_attachment=True,
                        download_name='GBBC_people_from_boa.csv',
                        mimetype='text/csv',
                    )
        return render_template('utilities_boa_people.html')

    @app.route('/utilities/boa-giving', methods=['GET', 'POST'])
    @require('utilities')
    def utilities_boa_giving():
        if request.method == 'POST':
            file = request.files.get('file')
            if not file or not file.filename.lower().endswith('.csv'):
                flash('Choose a Bank of America CSV file.', 'error')
            else:
                try:
                    gifts = extract_gifts_from_boa_csv(file)
                except ValueError as exc:
                    flash(str(exc), 'error')
                else:
                    output = io.StringIO()
                    writer = csv.DictWriter(output, fieldnames=[
                        'date', 'first_name', 'last_name', 'email', 'fund', 'amount',
                        'method', 'fiscal_year', 'reference', 'note',
                    ])
                    writer.writeheader()
                    writer.writerows(gifts)
                    content = io.BytesIO(output.getvalue().encode('utf-8'))
                    audit('export', 'utility', summary=f'BoA Zelle giving list ({len(gifts)})')
                    db.session.commit()
                    return send_file(
                        content,
                        as_attachment=True,
                        download_name='GBBC_contributions_from_boa_general_fund.csv',
                        mimetype='text/csv',
                    )
        return render_template('utilities_boa_giving.html')

    @app.route('/help', methods=['GET', 'POST'])
    def help_page():
        if not g.user:
            return redirect(url_for('login', next=request.path))
        if request.method == 'POST':
            subject = request.form.get('subject', '').strip(); message = request.form.get('message', '').strip()
            if not subject or not message: flash('Subject and description are required.', 'error')
            else:
                feedback = Feedback(user_id=g.user.id, category=request.form.get('category', 'Suggestion'), subject=subject, message=message)
                db.session.add(feedback); db.session.flush(); audit('create', 'feedback', feedback.id); db.session.commit()
                flash('Feedback sent.', 'success'); return redirect(url_for('help_page'))
        return render_template('help.html')

    @app.route('/approvals')
    @require('users')
    def approvals_redirect():
        """Legacy User Management URL now lives under Administration → Users."""
        return redirect(url_for('users'))

    @app.route('/admin/feedback', methods=['GET', 'POST'])
    @require('feedback_review')
    def feedback_review():
        if request.method == 'POST':
            item = db.get_or_404(Feedback, request.form.get('feedback_id', type=int))
            status = request.form.get('status', '')
            if status not in ('new', 'under review', 'resolved', 'declined'): abort(400)
            try:
                reason = require_audit_reason()
            except ValueError as exc:
                flash(str(exc), 'error')
                return redirect(url_for('feedback_review'))
            old = item.status
            item.status = status
            audit(
                'status', 'feedback', item.id, summary=item.subject, reason=reason,
                changes=[('Status', old, status)],
            )
            db.session.commit()
            return redirect(url_for('feedback_review'))
        items = db.session.scalars(db.select(Feedback).order_by(Feedback.created_at.desc()).limit(200)).all()
        return render_template('feedback_review.html', items=items)

    @app.route('/admin/backup', methods=['GET', 'POST'])
    @require('backup')
    def backup_page():
        years = db.session.scalars(db.select(FiscalYear).order_by(FiscalYear.start_date.desc())).all()
        selected_year_id = request.values.get('year_id', type=int)
        selected_year = db.session.get(FiscalYear, selected_year_id) if selected_year_id else (years[0] if years else None)
        pending = session.get('backup_restore')
        pending_summary = pending.get('summary') if pending and pending.get('user_id') == g.user.id else None

        if request.method == 'POST':
            action = (request.form.get('action') or '').strip()
            year = db.session.get(FiscalYear, request.form.get('year_id', type=int))
            if action == 'download':
                if not year:
                    flash('Choose a fiscal year to back up.', 'error')
                    return redirect(url_for('backup_page'))
                wb, summary = build_year_backup_workbook(year)
                out = io.BytesIO()
                wb.save(out)
                out.seek(0)
                audit(
                    'backup',
                    'fiscal_year',
                    year.id,
                    summary=f'Excel backup for {year.name}',
                    details=(
                        f'People: {summary["people"]}\n'
                        f'Contributions: {summary["contributions"]}\n'
                        f'Users: {summary["users"]}\n'
                        f'Audit log: {summary["audit"]}'
                    ),
                )
                db.session.commit()
                stamp = utcnow().strftime('%Y%m%d_%H%M')
                return send_file(
                    out,
                    as_attachment=True,
                    download_name=f'GBBC_backup_{year.name}_{stamp}.xlsx',
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                )

            if action == 'upload':
                if not year:
                    flash('Choose the fiscal year this backup should restore into.', 'error')
                    return redirect(url_for('backup_page'))
                upload = request.files.get('backup_file')
                if not upload or not upload.filename:
                    flash('Choose a backup Excel file to restore.', 'error')
                    return redirect(url_for('backup_page', year_id=year.id))
                if not upload.filename.lower().endswith('.xlsx'):
                    flash('Backup file must be an .xlsx workbook.', 'error')
                    return redirect(url_for('backup_page', year_id=year.id))
                raw = upload.read()
                if not raw:
                    flash('The uploaded backup file was empty.', 'error')
                    return redirect(url_for('backup_page', year_id=year.id))
                try:
                    parsed = parse_backup_workbook(io.BytesIO(raw), year)
                except Exception as exc:
                    flash(str(exc), 'error')
                    return redirect(url_for('backup_page', year_id=year.id))
                pending_path = backup_pending_path(g.user.id)
                pending_path.write_bytes(raw)
                token = secrets.token_urlsafe(16)
                session['backup_restore'] = {
                    'user_id': g.user.id,
                    'year_id': year.id,
                    'token': token,
                    'summary': parsed['summary'],
                    'path': str(pending_path),
                }
                return redirect(url_for('backup_page', year_id=year.id, confirm=1))

            if action == 'confirm_restore':
                pending = session.get('backup_restore') or {}
                token = (request.form.get('confirm_token') or '').strip()
                typed = (request.form.get('confirm_text') or '').strip().upper()
                if (
                    not pending
                    or pending.get('user_id') != g.user.id
                    or pending.get('token') != token
                    or typed != 'RESTORE'
                ):
                    flash('Restore confirmation failed. Upload the backup again and type RESTORE to confirm.', 'error')
                    session.pop('backup_restore', None)
                    return redirect(url_for('backup_page'))
                year = db.session.get(FiscalYear, pending.get('year_id'))
                path = Path(pending.get('path') or '')
                if not year or not path.exists():
                    flash('The staged backup file is missing. Upload it again.', 'error')
                    session.pop('backup_restore', None)
                    return redirect(url_for('backup_page'))
                try:
                    parsed = parse_backup_workbook(path, year)
                    apply_year_backup_restore(parsed, year, g.user)
                    audit(
                        'restore',
                        'fiscal_year',
                        year.id,
                        summary=f'Restored Excel backup for {year.name}',
                        reason='Confirmed restore from Administration backup page',
                        details=(
                            f'People upserted: {parsed["summary"]["people"]}\n'
                            f'Contributions replaced: {parsed["summary"]["contributions"]}\n'
                            f'Users upserted: {parsed["summary"]["users"]}\n'
                            f'Audit rows replaced: {parsed["summary"]["audit"]}'
                        ),
                    )
                    db.session.commit()
                except Exception as exc:
                    db.session.rollback()
                    current_app.logger.exception('Backup restore failed')
                    flash(f'Restore failed: {exc}', 'error')
                    return redirect(url_for('backup_page', year_id=year.id if year else None))
                finally:
                    session.pop('backup_restore', None)
                    try:
                        path.unlink(missing_ok=True)
                    except OSError:
                        pass
                flash(f'Restored backup for fiscal year {year.name}.', 'success')
                return redirect(url_for('backup_page', year_id=year.id))

            if action == 'cancel_restore':
                pending = session.pop('backup_restore', None) or {}
                path = Path(pending.get('path') or '')
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    pass
                flash('Restore cancelled.', 'success')
                return redirect(url_for('backup_page', year_id=pending.get('year_id')))

            flash('Unknown backup action.', 'error')
            return redirect(url_for('backup_page'))

        return render_template(
            'backup.html',
            years=years,
            selected_year=selected_year,
            pending_summary=pending_summary,
            pending_token=(pending or {}).get('token') if pending_summary else None,
            show_confirm=bool(request.args.get('confirm') and pending_summary),
        )

    @app.route('/admin/audit')
    @require('audit_view')
    def audit_view():
        events = db.session.scalars(db.select(Audit).order_by(Audit.timestamp.desc()).limit(200)).all()
        user_ids = {event.user_id for event in events if event.user_id}
        names = {user_id: db.session.get(User, user_id).name for user_id in user_ids if db.session.get(User, user_id)}
        return render_template('audit.html', events=events, names=names, format_eastern=format_eastern)

    @app.cli.command('init-db')
    def init_db():
        db.create_all()
        ensure_schema()
        if not db.session.scalar(db.select(Fund.id).limit(1)):
            db.session.add_all([Fund(name=n) for n in DEFAULT_FUNDS])
        if not db.session.scalar(db.select(FiscalYear.id).limit(1)):
            y = date.today().year; db.session.add(FiscalYear(name=str(y), start_date=date(y, 1, 1), end_date=date(y, 12, 31)))
        db.session.commit(); click.echo('Database initialized.')

    @app.cli.command('create-super-admin')
    @click.option('--email', prompt=True)
    @click.option('--first-name', prompt=True)
    @click.option('--last-name', prompt=True)
    @click.option('--phone', prompt=True)
    @click.password_option()
    def create_super_admin(email, first_name, last_name, phone, password):
        if not valid_email(email) or not all((first_name, last_name, phone)): raise click.ClickException('Provide a valid email, names, and phone.')
        if db.session.scalar(db.select(User).filter_by(email=email.lower())): raise click.ClickException('Email already exists.')
        user = User(email=email.lower(), first_name=first_name, last_name=last_name, phone=phone, role='Super Admin', password_hash=hash_password(password))
        db.session.add(user); db.session.commit(); click.echo('Super Admin created.')

    @app.cli.command('list-users')
    def list_users():
        users = db.session.scalars(db.select(User).order_by(User.last_name, User.first_name)).all()
        if not users: raise click.ClickException('No users exist. Run create-super-admin first.')
        for user in users:
            click.echo(f'{user.email}\t{user.name}\t{user.role}\t{"active" if user.active else "inactive"}')

    @app.cli.command('reset-password')
    @click.option('--email', prompt=True)
    @click.password_option()
    def reset_password(email, password):
        user = db.session.scalar(db.select(User).filter_by(email=email.strip().lower()))
        if not user: raise click.ClickException('No user has that email address.')
        if len(password) < 12: raise click.ClickException('A replacement password must contain at least 12 characters.')
        user.password_hash = hash_password(password)
        user.active = True
        db.session.execute(db.delete(LoginAttempt).where(LoginAttempt.email == user.email))
        db.session.add(Audit(user_id=user.id, action='password_reset_cli', entity='user', entity_id=user.id, summary='reset from command line'))
        db.session.commit()
        click.echo(f'Password reset for {user.email}.')

    return app

app = create_app()
